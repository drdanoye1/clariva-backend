"""
Engine 16 — Award & Project Management
Turns a won pipeline opportunity (Phase 4's FOARecord.pipeline_stage ==
"awarded") into an active Award: budget administration (burn-rate/variance
against the Budget Builder's BudgetRecord baseline), a terms/conditions
compliance checklist, amendments routed through Phase 3's generic
ApprovalRequest workflow, project execution status rolled up from Phase 2's
Scope of Work Engine, issue tracking, KPI performance actuals against
ProjectKnowledge.kpis targets, structured reports, closeout (feeding
Organizational Memory), and renewals (new Phase 4 pipeline entries)
(Clariva Enterprise™ PRD §16-17).

Design notes (same discipline as engines/scope_of_work_engine.py and
engines/collaboration_engine.py):
- Every method takes an already-open AsyncSession and only flushes, never
  commits — the caller's request-scoped session commits once, after the
  router handler returns.
- No permission checks live here — routers/awards.py gates access via
  workspace_access.py's proposal-access resolver (an Award is 1:1 with a
  Proposal, so "can edit this proposal" is exactly "can manage this award").
- Award creation is always explicit (a router action), never automatic on a
  pipeline-stage change — reaching "awarded" doesn't carry an award number,
  exact period of performance, or total value, so auto-creating an
  incomplete Award would be worse than requiring one deliberate step.
- Amendments reuse ApprovalRequest/Notification models directly (imported
  from db_models, not via CollaborationEngine) — the same
  "engines interoperate through shared models, not each other's classes"
  precedent set by funding_intelligence_engine.py's Notification usage.
  The actual decision — approve/reject, and applying `effective_changes` to
  the Award — happens inside collaboration_engine.py::decide_approval_request
  so there is exactly one place amendments get decided, whichever router the
  call came through.
- Project execution status is a live rollup read directly from Phase 2's
  WorkPackage/Task/Milestone/Deliverable/ProjectKnowledge tables — Phase 5
  does not duplicate a WBS or maintain its own copy.
- Every create/update path that returns a row for serialization calls
  `await db.refresh(obj)` after `await db.flush()` (server_default/onupdate
  timestamps are not eagerly fetched by AsyncSession + aiosqlite — see the
  identical fix first documented in credit_engine.py).
"""
from __future__ import annotations

import io
import json
import logging
import re
import traceback
from datetime import datetime
from typing import Any, Dict, List, Optional

import openai
from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

import storage
from config import settings
from models.db_models import (
    ApprovalRequest, Award, AwardAmendment, AwardCloseout, AwardComplianceItem,
    AwardCondition, AwardExpenditure, AwardPerformanceRecord, AwardReport, BudgetRecord,
    Deliverable, Document, DocumentVersion, FOARecord, MemoryRecord, Milestone,
    Notification, ProjectBaseline, ProjectIssue, ProjectKnowledge, Proposal,
    ScopeOfWork, StoredFile, Task, WorkPackage, new_uuid,
)
from models.schemas import AwardReportOut, BudgetStatusOut, PlannedVsActualOut, ProjectExecutionStatusOut

_log = logging.getLogger(__name__)


def _ai_error(exc: Exception) -> HTTPException:
    """See the identical helper in engines/scope_of_work_engine.py and
    engines/proposal_generator.py — this codebase's convention is a local
    copy per module rather than a shared import."""
    if isinstance(exc, openai.RateLimitError):
        _log.error("AI service rate limit exceeded: %s", exc)
        return HTTPException(status_code=503, detail="The AI writing service is temporarily unavailable due to high demand. Please try again in a few minutes.")
    if isinstance(exc, openai.AuthenticationError):
        _log.critical("AI service authentication failure: %s", exc)
        return HTTPException(status_code=503, detail="The AI writing service is not configured correctly. Please contact support.")
    if isinstance(exc, openai.APIConnectionError):
        _log.error("AI service connection error: %s", exc)
        return HTTPException(status_code=503, detail="The AI writing service is unreachable. Please try again in a moment.")
    if isinstance(exc, openai.APIStatusError):
        _log.error("AI service API error %s: %s", getattr(exc, "status_code", "?"), getattr(exc, "message", str(exc)))
        return HTTPException(status_code=503, detail="The AI writing service returned an unexpected error. Please try again.")
    _log.error("Unexpected report-narrative generation error:\n%s", traceback.format_exc())
    return HTTPException(status_code=500, detail="Content generation failed. Please try again or contact support.")


def _parse_json_response(raw: str) -> Dict[str, Any]:
    """Strip markdown fences and parse JSON, with a brace-scan fallback —
    mirrors engines/scope_of_work_engine.py's helper of the same name
    exactly (this codebase's convention is a local copy per module rather
    than a shared import — see that module's near-identical docstring)."""
    cleaned = re.sub(r"^```(?:json)?\s*", "", raw.strip())
    cleaned = re.sub(r"\s*```$", "", cleaned)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        start, end = cleaned.find("{"), cleaned.rfind("}") + 1
        try:
            return json.loads(cleaned[start:end])
        except Exception:
            raise HTTPException(status_code=500, detail="Could not parse the AI-extracted award information.")


def _naive(dt: Optional[datetime]) -> Optional[datetime]:
    """Strip tzinfo so period-of-performance math works whether the value
    came back tz-aware (PostgreSQL) or naive (SQLite)."""
    if dt is not None and dt.tzinfo is not None:
        return dt.replace(tzinfo=None)
    return dt


def _build_report_docx(buffer: "io.BytesIO", award: Award, report: "AwardReport") -> None:
    """Renders an *approved* AwardReport to a DOCX buffer — reuses
    utils/doc_utils.py's federal-style helpers (same look as every other
    exported document in this codebase) rather than a one-off style, but
    is otherwise deliberately simple: a report is a narrative plus a
    handful of frozen stats, not a multi-section proposal with figures/
    TOC, so it doesn't need document_output.py's heavier machinery."""
    from docx import Document as DocxDocument
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from utils.doc_utils import (
        FEDERAL_FONT, H1_PT, add_body_para, add_federal_heading,
        apply_federal_margins, _make_run, _para_spacing,
    )

    doc = DocxDocument()
    apply_federal_margins(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _para_spacing(title, after=4, before=0)
    _make_run(title, f"{report.report_type.replace('_', ' ').title()} Report", bold=True, pt=H1_PT, font=FEDERAL_FONT)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _para_spacing(meta, after=12)
    _make_run(meta, f"Award: {award.award_number or award.funding_agency}  |  Approved and exported {datetime.utcnow().strftime('%Y-%m-%d')}", pt=10, font=FEDERAL_FONT)

    add_federal_heading(doc, "Narrative", level=2)
    for paragraph in report.narrative.split("\n"):
        add_body_para(doc, paragraph)

    data = report.report_data or {}
    budget_status = data.get("budget_status") or {}
    execution_status = data.get("execution_status") or {}
    add_federal_heading(doc, "Summary Statistics (as of report generation)", level=2)
    stats_lines = [
        f"Burn rate: {budget_status.get('burn_rate_pct', 'N/A')}%",
        f"Work packages: {execution_status.get('total_work_packages', 'N/A')}",
        f"Open issues: {execution_status.get('open_issues', 'N/A')}",
        f"KPI records: {len(data.get('performance') or [])}",
    ]
    for line in stats_lines:
        add_body_para(doc, line)

    doc.save(buffer)


class AwardEngine:
    def __init__(self):
        self.client = openai.AsyncOpenAI(api_key=settings.OPENAI_API_KEY)

    async def _notify(self, db: AsyncSession, user_id: str, type_: str, message: str,
                       object_type: Optional[str] = None, object_id: Optional[str] = None) -> Notification:
        n = Notification(id=new_uuid(), user_id=user_id, type=type_, message=message,
                          object_type=object_type, object_id=object_id)
        db.add(n)
        await db.flush()
        return n

    # ── Awards ───────────────────────────────────────────────────────────────

    async def get_award_or_404(self, db: AsyncSession, award_id: str) -> Award:
        result = await db.execute(select(Award).where(Award.id == award_id))
        award = result.scalar_one_or_none()
        if not award:
            raise HTTPException(status_code=404, detail="Award not found")
        return award

    async def get_award_by_proposal(self, db: AsyncSession, proposal_id: str) -> Optional[Award]:
        result = await db.execute(select(Award).where(Award.proposal_id == proposal_id))
        return result.scalar_one_or_none()

    async def create_award(
        self, db: AsyncSession, proposal_id: str, data: Dict[str, Any],
        created_by: str, org_id: Optional[str],
    ) -> Award:
        existing = await self.get_award_by_proposal(db, proposal_id)
        if existing:
            raise HTTPException(status_code=400, detail="This proposal already has an award.")

        budget_record_id = None
        if data.get("link_budget", True):
            budget_result = await db.execute(select(BudgetRecord).where(BudgetRecord.proposal_id == proposal_id))
            budget = budget_result.scalar_one_or_none()
            if budget:
                budget_record_id = budget.id

        foa_id = None
        proposal_result = await db.execute(select(Proposal).where(Proposal.id == proposal_id))
        proposal = proposal_result.scalar_one_or_none()
        if proposal and proposal.foa_id:
            foa_id = proposal.foa_id

        award = Award(
            id=new_uuid(), proposal_id=proposal_id, foa_id=foa_id, org_id=org_id,
            budget_record_id=budget_record_id,
            award_number=data.get("award_number"), funding_agency=data["funding_agency"],
            period_of_performance_start=data.get("period_of_performance_start"),
            period_of_performance_end=data.get("period_of_performance_end"),
            total_award_value=data.get("total_award_value"), terms=data.get("terms"),
            status="active", created_by=created_by,
        )
        db.add(award)
        await db.flush()
        await db.refresh(award)
        return award

    # ── Quick Award Intake (Version 3.0 upgrade, Phase 15) ──────────────────
    # For a customer who already has a signed/funded award and never used
    # Pre-Award — there is no existing Proposal for create_award() above to
    # attach to (Award.proposal_id is unique/required — see that column's
    # docstring). Rather than change that relationship, this auto-creates a
    # minimal Proposal shell (origin="imported", no ProposalSection scaffold,
    # not meant to be opened in the proposal editor) purely so the Award has
    # somewhere to point, then reuses create_award() completely unchanged.

    async def create_award_from_intake(self, db: AsyncSession, data: Dict[str, Any], created_by: str) -> Award:
        shell = Proposal(
            id=new_uuid(), owner_id=created_by, title=data["title"],
            agency=data["funding_agency"], phase="phase_i", grant_type="federal_other",
            status="submitted", origin="imported",
        )
        db.add(shell)
        await db.flush()
        await db.refresh(shell)

        award = await self.create_award(
            db, shell.id,
            {k: v for k, v in data.items() if k not in ("title", "org_id")},
            created_by=created_by, org_id=data.get("org_id"),
        )
        return award

    async def update_award(self, db: AsyncSession, award_id: str, data: Dict[str, Any]) -> Award:
        award = await self.get_award_or_404(db, award_id)
        for field in ("award_number", "period_of_performance_start", "period_of_performance_end",
                      "total_award_value", "terms", "status"):
            if field in data and data[field] is not None:
                setattr(award, field, data[field])
        await db.flush()
        await db.refresh(award)
        return award

    async def list_awards_for_user(self, db: AsyncSession, user_id: str, org_ids: List[str]) -> List[Award]:
        owned = await db.execute(
            select(Award).join(Proposal, Proposal.id == Award.proposal_id).where(Proposal.owner_id == user_id)
        )
        by_id: Dict[str, Award] = {a.id: a for a in owned.scalars().all()}
        if org_ids:
            shared = await db.execute(select(Award).where(Award.org_id.in_(org_ids)))
            for a in shared.scalars().all():
                by_id[a.id] = a
        return sorted(by_id.values(), key=lambda a: a.created_at or datetime.min, reverse=True)

    # ── Award Received: activation & baselining (Phase 7 data model; PRD  ───
    # addendum "Version 3.0 — Three Synchronized Operational Environments").
    # `_build_snapshot` is the one place that reads the live BudgetRecord/
    # ScopeOfWork hierarchy into a denormalized snapshot dict, shared by
    # `activate_award` (the very first baseline) and `create_baseline_version`
    # (every re-baseline after an approved amendment), so the two never drift.

    async def _build_snapshot(self, db: AsyncSession, award: Award) -> Dict[str, Dict[str, Any]]:
        budget_snapshot: Dict[str, Any] = {}
        if award.budget_record_id:
            result = await db.execute(select(BudgetRecord).where(BudgetRecord.id == award.budget_record_id))
            budget = result.scalar_one_or_none()
            if budget:
                budget_snapshot = {
                    "budget_months": budget.budget_months,
                    "personnel": budget.personnel or [],
                    "consultants": budget.consultants or [],
                    "equipment": budget.equipment or [],
                    "travel": budget.travel or [],
                    "other_direct": budget.other_direct or [],
                    "subcontracts": budget.subcontracts or [],
                    "indirect_rate": budget.indirect_rate,
                    "indirect_base": budget.indirect_base,
                    "fee_rate": budget.fee_rate,
                    "total_direct": budget.total_direct,
                    "total_indirect": budget.total_indirect,
                    "total_cost": budget.total_cost,
                }

        scope_snapshot: Dict[str, Any] = {"work_packages": [], "milestones": [], "deliverables": []}
        pk_result = await db.execute(select(ProjectKnowledge).where(ProjectKnowledge.proposal_id == award.proposal_id))
        pk = pk_result.scalar_one_or_none()
        if pk:
            sow_result = await db.execute(select(ScopeOfWork).where(ScopeOfWork.project_knowledge_id == pk.id))
            sow = sow_result.scalar_one_or_none()
            if sow:
                wp_result = await db.execute(
                    select(WorkPackage).where(WorkPackage.scope_of_work_id == sow.id).order_by(WorkPackage.order_index)
                )
                work_packages = list(wp_result.scalars().all())
                scope_snapshot["work_packages"] = [
                    {
                        "id": wp.id, "name": wp.name, "description": wp.description, "lead": wp.lead,
                        "start_month": wp.start_month, "end_month": wp.end_month, "estimated_cost": wp.estimated_cost,
                    }
                    for wp in work_packages
                ]
                wp_ids = [wp.id for wp in work_packages]
                if wp_ids:
                    tasks_result = await db.execute(select(Task).where(Task.work_package_id.in_(wp_ids)))
                    scope_snapshot["tasks"] = [
                        {
                            "id": t.id, "work_package_id": t.work_package_id, "name": t.name,
                            "start_month": t.start_month, "end_month": t.end_month, "status": t.status,
                        }
                        for t in tasks_result.scalars().all()
                    ]
                milestones_result = await db.execute(select(Milestone).where(Milestone.scope_of_work_id == sow.id))
                scope_snapshot["milestones"] = [
                    {"id": m.id, "name": m.name, "due_month": m.due_month, "status": m.status}
                    for m in milestones_result.scalars().all()
                ]
                deliverables_result = await db.execute(select(Deliverable).where(Deliverable.scope_of_work_id == sow.id))
                scope_snapshot["deliverables"] = [
                    {"id": d.id, "name": d.name, "due_month": d.due_month,
                     "deliverable_type": d.deliverable_type, "status": d.status}
                    for d in deliverables_result.scalars().all()
                ]

        return {"budget_snapshot": budget_snapshot, "scope_snapshot": scope_snapshot}

    async def activate_award(self, db: AsyncSession, award_id: str, data: Dict[str, Any], created_by: str) -> ProjectBaseline:
        """The 'Activate Project' action: locks the first ProjectBaseline and
        flips Award.award_status from 'received' to 'active'. Only valid
        exactly once per award — re-baselining after this point goes through
        create_baseline_version (driven by an approved AwardAmendment), never
        through this method again."""
        award = await self.get_award_or_404(db, award_id)
        if award.award_status != "received":
            raise HTTPException(status_code=400, detail=f"This award is already {award.award_status} and cannot be activated again.")

        snapshot = await self._build_snapshot(db, award)
        baseline = ProjectBaseline(
            id=new_uuid(), award_id=award_id, version=1, is_current=True,
            total_award_value=award.total_award_value,
            period_of_performance_start=award.period_of_performance_start,
            period_of_performance_end=award.period_of_performance_end,
            budget_snapshot=snapshot["budget_snapshot"], scope_snapshot=snapshot["scope_snapshot"],
            notes=data.get("notes"), created_by=created_by,
        )
        db.add(baseline)
        award.award_status = "active"
        await db.flush()
        await db.refresh(baseline)
        return baseline

    async def create_baseline_version(self, db: AsyncSession, award_id: str, data: Dict[str, Any], created_by: str) -> ProjectBaseline:
        """Re-baselines an already-active award — called after an
        AwardAmendment changing budget, scope, or schedule is approved
        (Phase 8 wires the actual trigger; this method is the mechanism).
        Never edits the previous baseline row: flips it to is_current=False
        and inserts a new, higher-version row instead, so every prior
        approved state stays reconstructable."""
        award = await self.get_award_or_404(db, award_id)
        current_result = await db.execute(
            select(ProjectBaseline).where(ProjectBaseline.award_id == award_id, ProjectBaseline.is_current == True)  # noqa: E712
        )
        current = current_result.scalar_one_or_none()
        if not current:
            raise HTTPException(status_code=400, detail="This award has not been activated yet — no baseline exists to re-version.")

        current.is_current = False
        await db.flush()

        snapshot = await self._build_snapshot(db, award)
        baseline = ProjectBaseline(
            id=new_uuid(), award_id=award_id, version=current.version + 1, is_current=True,
            total_award_value=award.total_award_value,
            period_of_performance_start=award.period_of_performance_start,
            period_of_performance_end=award.period_of_performance_end,
            budget_snapshot=snapshot["budget_snapshot"], scope_snapshot=snapshot["scope_snapshot"],
            notes=data.get("notes"), created_by=created_by,
        )
        db.add(baseline)
        await db.flush()
        await db.refresh(baseline)
        return baseline

    async def get_current_baseline(self, db: AsyncSession, award_id: str) -> Optional[ProjectBaseline]:
        result = await db.execute(
            select(ProjectBaseline).where(ProjectBaseline.award_id == award_id, ProjectBaseline.is_current == True)  # noqa: E712
        )
        return result.scalar_one_or_none()

    async def list_baselines(self, db: AsyncSession, award_id: str) -> List[ProjectBaseline]:
        result = await db.execute(
            select(ProjectBaseline).where(ProjectBaseline.award_id == award_id).order_by(ProjectBaseline.version.desc())
        )
        return list(result.scalars().all())

    # ── Award Intake Intelligence (Version 3.0 upgrade, Phase D) ────────────
    # Quick Award Intake (Phase 15) never touches Scope of Work or Budget —
    # a customer skipping Pre-Award has neither by design. That means
    # activate_award()'s baseline snapshot above (_build_snapshot) captures
    # nothing, and Planned vs. Actual/Budget Status stay permanently blank,
    # even though the customer's uploaded award notice/funded proposal often
    # already states the award value, dates, and sometimes a work plan. This
    # reads that already-uploaded document TEXT (Document Library, populated
    # by routers/awards.py::intake_award_document) and proposes structured
    # data to fill the gap — reviewed and applied explicitly via
    # POST /awards/{award_id}/apply-intelligence, never silently.

    async def gather_intake_document_text(self, db: AsyncSession, proposal_id: str) -> str:
        """Concatenates the extracted text of every Quick Award Intake
        document attached to this proposal (funded_proposal + award_notice
        library types). Uses each Document's highest-version_number
        DocumentVersion — normally there's exactly one version per intake
        upload, but this stays correct if a document is ever replaced."""
        docs_result = await db.execute(
            select(Document).where(
                Document.proposal_id == proposal_id,
                Document.library_type.in_(["funded_proposal", "award_notice"]),
                Document.status == "active",
            )
        )
        blocks: List[str] = []
        for doc in docs_result.scalars().all():
            v_result = await db.execute(
                select(DocumentVersion).where(DocumentVersion.document_id == doc.id)
                .order_by(DocumentVersion.version_number.desc())
            )
            latest = v_result.scalars().first()
            if latest and (latest.content or "").strip():
                blocks.append(f"### {doc.title}\n{latest.content.strip()}")
        return "\n\n".join(blocks)

    async def extract_scope_and_award_fields(self, proposal: Any, document_text: str) -> Dict[str, Any]:
        """AI extraction of award value, period of performance, a work
        breakdown, and Project Knowledge fields from the customer's uploaded
        document text. work_packages mirrors scope_of_work_engine.py's
        generate_work_breakdown() shape exactly, so it can be persisted via
        that engine's apply_generated_work_breakdown() unchanged;
        objectives/need_statement/outputs/outcomes/kpis mirror
        ProjectKnowledge's own columns (see models/db_models.py) and
        scope_of_work_engine.py's derive_project_knowledge_from_proposal()'s
        field set — the same summary that feature derives from AI-*generated*
        ProposalSection content, derived here instead from the customer's
        *uploaded* documents, since a Quick-Award-Intake or normal-Award-page
        upload has no generated sections to derive from (see
        ScopeOfWorkEngine.derive_project_knowledge_from_proposal's 400 when a
        proposal has none — this method exists so that path isn't a dead end
        for these proposals). Returns an unpersisted dict for the caller to
        show for review — nothing is saved here.

        Truncates document_text at 100k chars, not the ~16k this originally
        shipped with — a real funded proposal easily runs 30-50k+ characters,
        and the Work Plan / Implementation Schedule section (where work
        packages/milestones actually live) is typically well past the
        halfway point, after Executive Summary/Statement of Need/Goals/
        Technical Approach. The old 16k cutoff silently dropped that section
        for any real-world document, so the AI correctly reported "no work
        plan found" — it never saw one. settings.OPENAI_MODEL is gpt-4o
        (128k-token context window, ~4 chars/token), so 100k chars (~25k
        tokens) leaves enormous headroom; this isn't a case of "raise it a
        little," the original limit was just wrong for this model."""
        prompt = f"""
Read the following award-related document(s) and extract structured project
information as JSON only (no markdown fences, no commentary — just the JSON
object).

Project title: {getattr(proposal, "title", "")}
Funding agency: {getattr(proposal, "agency", "")}

--- DOCUMENT(S) ---
{document_text[:100000]}
--- END DOCUMENT(S) ---

Return JSON matching exactly this shape:
{{
  "total_award_value": number or null,
  "period_of_performance_start": "YYYY-MM-DD" or null,
  "period_of_performance_end": "YYYY-MM-DD" or null,
  "work_packages": [
    {{
      "name": "string",
      "description": "string",
      "start_month": 1,
      "end_month": 6,
      "tasks": ["string", "..."],
      "milestones": ["string", "..."],
      "deliverables": ["string", "..."]
    }}
  ],
  "objectives": "string or null",
  "need_statement": "string or null",
  "outputs": "string or null",
  "outcomes": "string or null",
  "kpis": [
    {{"name": "string", "target": "string", "unit": "string"}}
  ],
  "extraction_notes": "string"
}}

Only extract what the document(s) actually state. Use null for
total_award_value/dates/objectives/need_statement/outputs/outcomes if not
stated, and empty arrays for work_packages/kpis if the document doesn't
describe a work plan or performance measures — a short award notice letter
often only states the amount and dates with nothing else, and that's fine;
don't invent any of it. Month numbers in work packages are 1-indexed from
project start.

For objectives/need_statement/outputs/outcomes, write plain prose (no
markdown, no bullets) summarizing what the document(s) actually say —
objectives is what the project aims to achieve, need_statement is the
problem/gap it addresses, outputs are concrete things produced, outcomes
are longer-term changes/results. For kpis, extract each named performance
indicator with its stated numeric or quantitative target and unit exactly
as written (e.g. a table row "Participants Trained: 250" becomes
{{"name": "Participants Trained", "target": "250", "unit": "participants"}});
leave target/unit as empty strings if the document names an indicator
without a stated target. For extraction_notes, write a specific one- or
two-sentence summary referencing what this particular document actually
contains (cite the title or a phrase from it) — not a generic template
sentence.
""".strip()
        try:
            response = await self.client.chat.completions.create(
                model=settings.OPENAI_MODEL,
                messages=[
                    {"role": "system", "content": "You are an expert grant analyst extracting structured award information from funder documents. Respond with a single JSON object only. Never invent data the document doesn't support."},
                    {"role": "user", "content": prompt},
                ],
                temperature=0.2,
                max_tokens=2000,
            )
        except Exception as exc:
            raise _ai_error(exc)
        return _parse_json_response(response.choices[0].message.content or "")

    # ── Award Received: sponsor conditions (Phase 7 data model) ────────────

    async def get_condition_or_404(self, db: AsyncSession, condition_id: str) -> AwardCondition:
        result = await db.execute(select(AwardCondition).where(AwardCondition.id == condition_id))
        condition = result.scalar_one_or_none()
        if not condition:
            raise HTTPException(status_code=404, detail="Award condition not found")
        return condition

    async def create_condition(self, db: AsyncSession, award_id: str, data: Dict[str, Any], created_by: str) -> AwardCondition:
        await self.get_award_or_404(db, award_id)
        condition = AwardCondition(
            id=new_uuid(), award_id=award_id, description=data["description"],
            category=data.get("category"), due_date=data.get("due_date"), created_by=created_by,
        )
        db.add(condition)
        await db.flush()
        await db.refresh(condition)
        return condition

    async def list_conditions(self, db: AsyncSession, award_id: str) -> List[AwardCondition]:
        result = await db.execute(
            select(AwardCondition).where(AwardCondition.award_id == award_id).order_by(AwardCondition.due_date)
        )
        return list(result.scalars().all())

    async def update_condition(self, db: AsyncSession, condition_id: str, data: Dict[str, Any], resolved_by: Optional[str] = None) -> AwardCondition:
        condition = await self.get_condition_or_404(db, condition_id)
        for field in ("description", "category", "due_date", "status"):
            if field in data and data[field] is not None:
                setattr(condition, field, data[field])
        if data.get("status") in ("resolved", "waived") and not condition.resolved_at:
            condition.resolved_at = datetime.utcnow()
            condition.resolved_by = resolved_by
        await db.flush()
        await db.refresh(condition)
        return condition

    # ── Planned vs. Actual (Version 3.0 upgrade, Phase 10) ──────────────────
    # Replaces the dollars-vs-time-only get_budget_status comparison with a
    # real one against the locked ProjectBaseline, across budget, work
    # packages, milestones, deliverables, and schedule.

    async def get_planned_vs_actual(self, db: AsyncSession, award_id: str) -> PlannedVsActualOut:
        award = await self.get_award_or_404(db, award_id)
        baseline = await self.get_current_baseline(db, award_id)
        if not baseline:
            raise HTTPException(status_code=400, detail="This award has not been activated yet — no baseline exists to compare against.")

        # Budget: baseline_snapshot's frozen total, not the live BudgetRecord
        # (get_budget_status's baseline) — see PlannedVsActualOut's docstring
        # for why these two are deliberately allowed to diverge.
        baseline_total_cost = (baseline.budget_snapshot or {}).get("total_cost") or 0.0
        expenditures = await self.list_expenditures(db, award_id)
        total_expended = sum(e.amount for e in expenditures)
        burn_rate_pct = round(total_expended / baseline_total_cost * 100, 1) if baseline_total_cost > 0 else None

        start = _naive(baseline.period_of_performance_start)
        end = _naive(baseline.period_of_performance_end)
        elapsed_pct: Optional[float] = None
        elapsed_months: Optional[float] = None
        if start and end and end > start:
            total_seconds = (end - start).total_seconds()
            elapsed_seconds = max(0.0, min(total_seconds, (datetime.utcnow() - start).total_seconds()))
            elapsed_pct = round(elapsed_seconds / total_seconds * 100, 1)
            total_months = total_seconds / (30.44 * 86400)
            elapsed_months = elapsed_pct / 100 * total_months

        budget_variance_pct = round(burn_rate_pct - elapsed_pct, 1) if burn_rate_pct is not None and elapsed_pct is not None else None

        # Scope: planned counts come from the frozen baseline snapshot;
        # current/completed/behind-schedule come from LIVE ScopeOfWork rows
        # (unlike the budget comparison, live status fields — not baseline
        # snapshot values, which never change — are what tell us what's
        # actually been completed since the baseline was locked).
        scope_snapshot = baseline.scope_snapshot or {}
        planned_work_packages = len(scope_snapshot.get("work_packages") or [])
        planned_milestones = len(scope_snapshot.get("milestones") or [])
        planned_deliverables = len(scope_snapshot.get("deliverables") or [])

        current_work_packages = 0
        current_milestones = milestones_completed = milestones_behind_schedule = 0
        current_deliverables = deliverables_completed = deliverables_behind_schedule = 0

        pk_result = await db.execute(select(ProjectKnowledge).where(ProjectKnowledge.proposal_id == award.proposal_id))
        pk = pk_result.scalar_one_or_none()
        if pk:
            sow_result = await db.execute(select(ScopeOfWork).where(ScopeOfWork.project_knowledge_id == pk.id))
            sow = sow_result.scalar_one_or_none()
            if sow:
                wp_result = await db.execute(select(WorkPackage).where(WorkPackage.scope_of_work_id == sow.id))
                current_work_packages = len(list(wp_result.scalars().all()))

                milestones_result = await db.execute(select(Milestone).where(Milestone.scope_of_work_id == sow.id))
                live_milestones = list(milestones_result.scalars().all())
                current_milestones = len(live_milestones)
                milestones_completed = sum(1 for m in live_milestones if m.status == "complete")
                if elapsed_months is not None:
                    milestones_behind_schedule = sum(
                        1 for m in live_milestones
                        if m.status != "complete" and m.due_month is not None and m.due_month <= elapsed_months
                    )

                deliverables_result = await db.execute(select(Deliverable).where(Deliverable.scope_of_work_id == sow.id))
                live_deliverables = list(deliverables_result.scalars().all())
                current_deliverables = len(live_deliverables)
                deliverables_completed = sum(1 for d in live_deliverables if d.status == "complete")
                if elapsed_months is not None:
                    deliverables_behind_schedule = sum(
                        1 for d in live_deliverables
                        if d.status != "complete" and d.due_month is not None and d.due_month <= elapsed_months
                    )

        scope_drift = (
            planned_work_packages != current_work_packages
            or planned_milestones != current_milestones
            or planned_deliverables != current_deliverables
        )

        return PlannedVsActualOut(
            award_id=award_id, baseline_version=baseline.version, baseline_locked_at=baseline.created_at,
            baseline_total_cost=baseline_total_cost, total_expended=total_expended,
            burn_rate_pct=burn_rate_pct, elapsed_pct=elapsed_pct, budget_variance_pct=budget_variance_pct,
            planned_work_packages=planned_work_packages, current_work_packages=current_work_packages,
            planned_milestones=planned_milestones, current_milestones=current_milestones,
            milestones_completed=milestones_completed, milestones_behind_schedule=milestones_behind_schedule,
            planned_deliverables=planned_deliverables, current_deliverables=current_deliverables,
            deliverables_completed=deliverables_completed, deliverables_behind_schedule=deliverables_behind_schedule,
            scope_drift=scope_drift,
        )

    # ── Budget administration / burn-rate (PRD §16) ─────────────────────────

    async def create_expenditure(self, db: AsyncSession, award_id: str, data: Dict[str, Any], recorded_by: str) -> AwardExpenditure:
        await self.get_award_or_404(db, award_id)
        exp = AwardExpenditure(
            id=new_uuid(), award_id=award_id, category=data["category"],
            description=data.get("description"), amount=data["amount"],
            incurred_date=data.get("incurred_date"), recorded_by=recorded_by,
        )
        db.add(exp)
        await db.flush()
        await db.refresh(exp)
        return exp

    async def list_expenditures(self, db: AsyncSession, award_id: str) -> List[AwardExpenditure]:
        result = await db.execute(
            select(AwardExpenditure).where(AwardExpenditure.award_id == award_id).order_by(AwardExpenditure.created_at.desc())
        )
        return list(result.scalars().all())

    async def get_budget_status(self, db: AsyncSession, award_id: str) -> BudgetStatusOut:
        award = await self.get_award_or_404(db, award_id)
        baseline_total_cost = 0.0
        if award.budget_record_id:
            budget_result = await db.execute(select(BudgetRecord).where(BudgetRecord.id == award.budget_record_id))
            budget = budget_result.scalar_one_or_none()
            if budget:
                baseline_total_cost = budget.total_cost or 0.0

        expenditures = await self.list_expenditures(db, award_id)
        total_expended = sum(e.amount for e in expenditures)
        by_category: Dict[str, float] = {}
        for e in expenditures:
            by_category[e.category] = by_category.get(e.category, 0.0) + e.amount

        burn_rate_pct = round(total_expended / baseline_total_cost * 100, 1) if baseline_total_cost > 0 else None

        elapsed_pct = None
        start = _naive(award.period_of_performance_start)
        end = _naive(award.period_of_performance_end)
        if start and end and end > start:
            total_seconds = (end - start).total_seconds()
            elapsed_seconds = max(0.0, min(total_seconds, (datetime.utcnow() - start).total_seconds()))
            elapsed_pct = round(elapsed_seconds / total_seconds * 100, 1)

        variance_pct = round(burn_rate_pct - elapsed_pct, 1) if burn_rate_pct is not None and elapsed_pct is not None else None

        return BudgetStatusOut(
            award_id=award_id, baseline_total_cost=baseline_total_cost, total_expended=total_expended,
            burn_rate_pct=burn_rate_pct, elapsed_pct=elapsed_pct, variance_pct=variance_pct, by_category=by_category,
        )

    # ── Compliance checklist (PRD §16) ──────────────────────────────────────

    async def get_compliance_item_or_404(self, db: AsyncSession, item_id: str) -> AwardComplianceItem:
        result = await db.execute(select(AwardComplianceItem).where(AwardComplianceItem.id == item_id))
        item = result.scalar_one_or_none()
        if not item:
            raise HTTPException(status_code=404, detail="Compliance item not found")
        return item

    async def create_compliance_item(self, db: AsyncSession, award_id: str, data: Dict[str, Any]) -> AwardComplianceItem:
        await self.get_award_or_404(db, award_id)
        item = AwardComplianceItem(
            id=new_uuid(), award_id=award_id, obligation=data["obligation"],
            category=data.get("category"), due_date=data.get("due_date"), notes=data.get("notes"),
        )
        db.add(item)
        await db.flush()
        await db.refresh(item)
        return item

    async def list_compliance_items(self, db: AsyncSession, award_id: str) -> List[AwardComplianceItem]:
        result = await db.execute(
            select(AwardComplianceItem).where(AwardComplianceItem.award_id == award_id).order_by(AwardComplianceItem.due_date)
        )
        return list(result.scalars().all())

    async def update_compliance_item(self, db: AsyncSession, item_id: str, data: Dict[str, Any], completed_by: Optional[str] = None) -> AwardComplianceItem:
        item = await self.get_compliance_item_or_404(db, item_id)
        for field in ("obligation", "category", "due_date", "status", "notes"):
            if field in data and data[field] is not None:
                setattr(item, field, data[field])
        if data.get("status") == "complete" and not item.completed_at:
            item.completed_at = datetime.utcnow()
            item.completed_by = completed_by
        await db.flush()
        await db.refresh(item)
        return item

    async def delete_compliance_item(self, db: AsyncSession, item_id: str) -> None:
        item = await self.get_compliance_item_or_404(db, item_id)
        await db.delete(item)
        await db.flush()

    # ── Amendments (PRD §16 — approval routing via Phase 3's ApprovalRequest) ──

    async def get_amendment_or_404(self, db: AsyncSession, amendment_id: str) -> AwardAmendment:
        result = await db.execute(select(AwardAmendment).where(AwardAmendment.id == amendment_id))
        amendment = result.scalar_one_or_none()
        if not amendment:
            raise HTTPException(status_code=404, detail="Amendment not found")
        return amendment

    async def create_amendment(
        self, db: AsyncSession, award_id: str, data: Dict[str, Any], requested_by: str, org_id: Optional[str],
    ) -> AwardAmendment:
        await self.get_award_or_404(db, award_id)
        amendment = AwardAmendment(
            id=new_uuid(), award_id=award_id, amendment_type=data["amendment_type"],
            description=data["description"], effective_changes=data.get("effective_changes"),
            requested_by=requested_by, status="pending",
        )
        db.add(amendment)
        await db.flush()
        await db.refresh(amendment)

        approver_id = data.get("approver_id")
        approval = ApprovalRequest(
            id=new_uuid(), org_id=org_id, object_type="award_amendment", object_id=amendment.id,
            requested_by=requested_by, approver_id=approver_id, notes=data["description"],
        )
        db.add(approval)
        await db.flush()
        await db.refresh(approval)
        amendment.approval_request_id = approval.id
        await db.flush()
        await db.refresh(amendment)

        if approver_id:
            await self._notify(db, approver_id, "approval_requested", "An award amendment needs your decision.",
                                object_type="award_amendment", object_id=amendment.id)
        return amendment

    async def list_amendments(self, db: AsyncSession, award_id: str) -> List[AwardAmendment]:
        result = await db.execute(
            select(AwardAmendment).where(AwardAmendment.award_id == award_id).order_by(AwardAmendment.created_at.desc())
        )
        return list(result.scalars().all())

    # ── Issues (PRD §17) ─────────────────────────────────────────────────────

    async def get_issue_or_404(self, db: AsyncSession, issue_id: str) -> ProjectIssue:
        result = await db.execute(select(ProjectIssue).where(ProjectIssue.id == issue_id))
        issue = result.scalar_one_or_none()
        if not issue:
            raise HTTPException(status_code=404, detail="Issue not found")
        return issue

    async def create_issue(self, db: AsyncSession, award_id: str, data: Dict[str, Any], raised_by: str) -> ProjectIssue:
        await self.get_award_or_404(db, award_id)
        issue = ProjectIssue(
            id=new_uuid(), award_id=award_id, work_package_id=data.get("work_package_id"),
            title=data["title"], description=data.get("description"),
            severity=data.get("severity") or "medium", raised_by=raised_by,
        )
        db.add(issue)
        await db.flush()
        await db.refresh(issue)
        return issue

    async def list_issues(self, db: AsyncSession, award_id: str) -> List[ProjectIssue]:
        result = await db.execute(
            select(ProjectIssue).where(ProjectIssue.award_id == award_id).order_by(ProjectIssue.created_at.desc())
        )
        return list(result.scalars().all())

    async def update_issue(self, db: AsyncSession, issue_id: str, data: Dict[str, Any]) -> ProjectIssue:
        issue = await self.get_issue_or_404(db, issue_id)
        for field in ("title", "description", "severity", "status"):
            if field in data and data[field] is not None:
                setattr(issue, field, data[field])
        if data.get("status") == "resolved" and not issue.resolved_at:
            issue.resolved_at = datetime.utcnow()
        await db.flush()
        await db.refresh(issue)
        return issue

    # ── Project execution status (PRD §17 — read from Phase 2's SOW Engine) ──

    async def get_execution_status(self, db: AsyncSession, award_id: str) -> ProjectExecutionStatusOut:
        award = await self.get_award_or_404(db, award_id)

        pk_result = await db.execute(select(ProjectKnowledge).where(ProjectKnowledge.proposal_id == award.proposal_id))
        pk = pk_result.scalar_one_or_none()
        open_issues_result = await db.execute(
            select(ProjectIssue).where(ProjectIssue.award_id == award_id, ProjectIssue.status == "open")
        )
        open_issues = len(list(open_issues_result.scalars().all()))

        if not pk:
            return ProjectExecutionStatusOut(award_id=award_id, open_issues=open_issues)

        sow_result = await db.execute(select(ScopeOfWork).where(ScopeOfWork.project_knowledge_id == pk.id))
        sow = sow_result.scalar_one_or_none()
        if not sow:
            return ProjectExecutionStatusOut(award_id=award_id, open_issues=open_issues, stale_flags=pk.stale_flags or {})

        wp_result = await db.execute(select(WorkPackage).where(WorkPackage.scope_of_work_id == sow.id))
        work_packages = list(wp_result.scalars().all())
        wp_ids = [wp.id for wp in work_packages]

        tasks: List[Task] = []
        if wp_ids:
            tasks_result = await db.execute(select(Task).where(Task.work_package_id.in_(wp_ids)))
            tasks = list(tasks_result.scalars().all())

        milestones_result = await db.execute(select(Milestone).where(Milestone.scope_of_work_id == sow.id))
        milestones = list(milestones_result.scalars().all())

        deliverables_result = await db.execute(select(Deliverable).where(Deliverable.scope_of_work_id == sow.id))
        deliverables = list(deliverables_result.scalars().all())

        def _tally(items, attr="status") -> Dict[str, int]:
            counts: Dict[str, int] = {}
            for item in items:
                key = getattr(item, attr) or "unknown"
                counts[key] = counts.get(key, 0) + 1
            return counts

        return ProjectExecutionStatusOut(
            award_id=award_id,
            total_work_packages=len(work_packages),
            total_tasks=len(tasks), tasks_by_status=_tally(tasks),
            total_milestones=len(milestones), milestones_by_status=_tally(milestones),
            total_deliverables=len(deliverables), deliverables_by_status=_tally(deliverables),
            open_issues=open_issues, stale_flags=pk.stale_flags or {},
        )

    # ── Performance / KPI actuals (PRD §17) ─────────────────────────────────

    async def create_performance_record(self, db: AsyncSession, award_id: str, data: Dict[str, Any], recorded_by: Optional[str]) -> AwardPerformanceRecord:
        await self.get_award_or_404(db, award_id)
        record = AwardPerformanceRecord(
            id=new_uuid(), award_id=award_id, kpi_name=data["kpi_name"],
            target=data.get("target"), actual_value=data.get("actual_value"),
            unit=data.get("unit"), period_label=data.get("period_label"),
            notes=data.get("notes"), recorded_by=recorded_by,
        )
        db.add(record)
        await db.flush()
        await db.refresh(record)
        return record

    async def list_performance_records(self, db: AsyncSession, award_id: str) -> List[AwardPerformanceRecord]:
        result = await db.execute(
            select(AwardPerformanceRecord).where(AwardPerformanceRecord.award_id == award_id).order_by(AwardPerformanceRecord.created_at.desc())
        )
        return list(result.scalars().all())

    # ── Reports (PRD §17 — "generated as views, not authored from scratch") ──

    async def generate_report(self, db: AsyncSession, award_id: str, report_type: str, award_out: Any) -> AwardReportOut:
        budget_status = await self.get_budget_status(db, award_id)
        execution_status = await self.get_execution_status(db, award_id)
        performance = await self.list_performance_records(db, award_id)

        compliance_items = await self.list_compliance_items(db, award_id)
        compliance_summary: Dict[str, int] = {}
        for item in compliance_items:
            compliance_summary[item.status] = compliance_summary.get(item.status, 0) + 1

        from models.schemas import AwardPerformanceRecordOut
        return AwardReportOut(
            award_id=award_id, report_type=report_type, generated_at=datetime.utcnow(),
            award=award_out, budget_status=budget_status, execution_status=execution_status,
            performance=[AwardPerformanceRecordOut.model_validate(p) for p in performance],
            compliance_summary=compliance_summary or None,
        )

    async def generate_report_narrative(self, report: AwardReportOut, proposal: Proposal, additional_context: Optional[str] = None) -> str:
        exec_status = report.execution_status
        budget_status = report.budget_status
        prompt = f"""
Write a {report.report_type} report narrative (3-5 paragraphs, plain prose, no
markdown or headers) for the following funded project, suitable for
submission to the funding agency or for internal project review.

Project title: {getattr(proposal, "title", "")}
Agency: {report.award.funding_agency}
Award number: {report.award.award_number or "Not yet assigned"}
Total award value: {report.award.total_award_value if report.award.total_award_value is not None else "Not specified"}
Period of performance: {report.award.period_of_performance_start} to {report.award.period_of_performance_end}

Budget status: {f"{budget_status.burn_rate_pct}% of baseline budget expended ({budget_status.total_expended} of {budget_status.baseline_total_cost})" if budget_status else "Not available"}
Elapsed period of performance: {f"{budget_status.elapsed_pct}%" if budget_status and budget_status.elapsed_pct is not None else "Not available"}

Work packages: {exec_status.total_work_packages if exec_status else 0}
Tasks by status: {exec_status.tasks_by_status if exec_status else {}}
Milestones by status: {exec_status.milestones_by_status if exec_status else {}}
Deliverables by status: {exec_status.deliverables_by_status if exec_status else {}}
Open issues: {exec_status.open_issues if exec_status else 0}

KPI performance records: {[(p.kpi_name, p.target, p.actual_value, p.unit) for p in report.performance]}
Compliance obligation status counts: {report.compliance_summary or {}}
{f"Additional context: {additional_context}" if additional_context else ""}

If information needed to write a specific, credible narrative is missing,
write around it generally rather than inventing specifics — do not
fabricate figures, dates, or results not given above.
""".strip()
        try:
            response = await self.client.chat.completions.create(
                model=settings.OPENAI_MODEL,
                messages=[
                    {"role": "system", "content": "You are an expert grant/project manager producing a funder-facing project report. Output plain prose only — no markdown, no headers, no bullet characters."},
                    {"role": "user", "content": prompt},
                ],
                temperature=0.5,
                max_tokens=900,
            )
        except Exception as exc:
            raise _ai_error(exc)
        return response.choices[0].message.content.strip()

    # ── Persisted, human-reviewed reports (see models/db_models.py's
    # AwardReport docstring — the generate_report/generate_report_narrative
    # pair above stays as an ephemeral live-preview pair; everything below
    # is the actual draft → edit → submit → approve/reject → export
    # workflow, routed through Phase 3's ApprovalRequest exactly like
    # amendments above, with collaboration_engine.py::decide_approval_request
    # applying the decision.) ───────────────────────────────────────────────

    async def get_award_report_or_404(self, db: AsyncSession, report_id: str) -> AwardReport:
        result = await db.execute(select(AwardReport).where(AwardReport.id == report_id))
        report = result.scalar_one_or_none()
        if not report:
            raise HTTPException(status_code=404, detail="Report not found")
        return report

    async def create_report_draft(
        self, db: AsyncSession, award_id: str, report_type: str, additional_context: Optional[str],
        award_out: Any, requested_by: str,
    ) -> AwardReport:
        """Runs the same generate_report/generate_report_narrative pipeline
        the ephemeral preview endpoints use, but persists the result as a
        draft instead of just handing it back — this is step one of the
        human-in-the-loop workflow. Nothing produced here is exportable
        until a person has reviewed (optionally edited) this draft and it
        has been approved."""
        report_view = await self.generate_report(db, award_id, report_type, award_out)

        award = await self.get_award_or_404(db, award_id)
        proposal_result = await db.execute(select(Proposal).where(Proposal.id == award.proposal_id))
        proposal = proposal_result.scalar_one_or_none()
        narrative = await self.generate_report_narrative(report_view, proposal, additional_context)

        report_data = json.loads(report_view.model_dump_json(exclude={"narrative", "award"}))
        report = AwardReport(
            id=new_uuid(), award_id=award_id, report_type=report_type, status="draft",
            narrative=narrative, ai_generated_narrative=narrative, report_data=report_data,
            requested_by=requested_by,
        )
        db.add(report)
        await db.flush()
        await db.refresh(report)
        return report

    async def update_report_draft(self, db: AsyncSession, report_id: str, narrative: str) -> AwardReport:
        report = await self.get_award_report_or_404(db, report_id)
        if report.status != "draft":
            raise HTTPException(status_code=400, detail="Only a draft report's narrative can be edited — this one has already been submitted for approval.")
        report.narrative = narrative
        await db.flush()
        await db.refresh(report)
        return report

    async def list_award_reports(self, db: AsyncSession, award_id: str) -> List[AwardReport]:
        result = await db.execute(
            select(AwardReport).where(AwardReport.award_id == award_id).order_by(AwardReport.created_at.desc())
        )
        return list(result.scalars().all())

    async def submit_report_for_approval(
        self, db: AsyncSession, report_id: str, requested_by: str, approver_id: Optional[str],
        notes: Optional[str], org_id: Optional[str],
    ) -> AwardReport:
        report = await self.get_award_report_or_404(db, report_id)
        if report.status != "draft":
            raise HTTPException(status_code=400, detail="This report has already been submitted for approval.")

        approval = ApprovalRequest(
            id=new_uuid(), org_id=org_id, object_type="award_report", object_id=report.id,
            requested_by=requested_by, approver_id=approver_id, notes=notes,
        )
        db.add(approval)
        await db.flush()
        await db.refresh(approval)
        report.approval_request_id = approval.id
        report.status = "pending_approval"
        await db.flush()
        await db.refresh(report)

        if approver_id:
            await self._notify(db, approver_id, "approval_requested", "A post-award report needs your review and decision.",
                                object_type="award_report", object_id=report.id)
        return report

    async def export_report(self, db: AsyncSession, report_id: str, created_by: str, org_id: Optional[str]) -> Dict[str, Any]:
        """Human-in-the-loop gate: a report can only become a downloadable
        file once a person has approved it (see models/db_models.py's
        AwardReport docstring) — there is no path from AI generation
        straight to a file."""
        report = await self.get_award_report_or_404(db, report_id)
        if report.status != "approved":
            raise HTTPException(status_code=400, detail="This report must be reviewed and approved before it can be exported.")

        award = await self.get_award_or_404(db, report.award_id)
        buffer = io.BytesIO()
        _build_report_docx(buffer, award, report)
        content = buffer.getvalue()
        filename = f"{report.report_type}_report_{report.id[:8]}.docx"
        content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        storage_key = await storage.upload_file(org_id, "award_report_export", content, filename, content_type)
        stored_file = StoredFile(
            id=new_uuid(), org_id=org_id, object_type="award_report_export", object_id=report.id,
            storage_key=storage_key, original_filename=filename, content_type=content_type,
            size_bytes=len(content), checksum=storage.sha256_hex(content), created_by=created_by,
        )
        db.add(stored_file)
        await db.flush()
        await db.refresh(stored_file)
        report.exported_file_id = stored_file.id
        await db.flush()
        download_url = await storage.get_download_url(storage_key, filename=filename)
        return {"download_url": download_url, "file_size": len(content), "exported_at": datetime.utcnow()}

    # ── Closeout (PRD §17 — feeds Organizational Memory) ────────────────────

    async def close_award(self, db: AsyncSession, award_id: str, data: Dict[str, Any], closed_by: str) -> AwardCloseout:
        award = await self.get_award_or_404(db, award_id)

        result = await db.execute(select(AwardCloseout).where(AwardCloseout.award_id == award_id))
        closeout = result.scalar_one_or_none()
        if not closeout:
            closeout = AwardCloseout(id=new_uuid(), award_id=award_id)
            db.add(closeout)

        closeout.deliverables_reconciled = data.get("deliverables_reconciled", closeout.deliverables_reconciled)
        closeout.deliverables_notes = data.get("deliverables_notes") or closeout.deliverables_notes
        closeout.equipment_disposition = data.get("equipment_disposition") or closeout.equipment_disposition
        closeout.final_report_submitted = data.get("final_report_submitted", closeout.final_report_submitted)
        lessons = data.get("lessons_learned")
        if lessons is not None:
            closeout.lessons_learned = lessons

        memory = MemoryRecord(
            id=new_uuid(), org_id=closed_by, proposal_id=award.proposal_id,
            agency=award.funding_agency, outcome="funded",
            score=data.get("outcome_score"), lessons_learned=closeout.lessons_learned or [],
        )
        db.add(memory)
        await db.flush()
        await db.refresh(memory)

        closeout.memory_record_id = memory.id
        closeout.closed_by = closed_by
        closeout.closed_at = datetime.utcnow()
        award.status = "closed"

        await db.flush()
        await db.refresh(closeout)
        return closeout

    async def get_closeout(self, db: AsyncSession, award_id: str) -> Optional[AwardCloseout]:
        result = await db.execute(select(AwardCloseout).where(AwardCloseout.award_id == award_id))
        return result.scalar_one_or_none()

    # ── Renewal (PRD §17 — new pipeline entry linked to the originating award) ──

    async def create_renewal_opportunity(self, db: AsyncSession, award_id: str, data: Dict[str, Any], uploaded_by: str) -> FOARecord:
        award = await self.get_award_or_404(db, award_id)
        proposal_result = await db.execute(select(Proposal).where(Proposal.id == award.proposal_id))
        proposal = proposal_result.scalar_one_or_none()
        if not proposal:
            raise HTTPException(status_code=404, detail="Originating proposal not found")

        record = FOARecord(
            id=new_uuid(), agency=award.funding_agency,
            program_title=data.get("program_title") or f"{proposal.title} (Renewal)",
            phase=proposal.phase, grant_type=proposal.grant_type or "sbir",
            deadline=data.get("deadline"), uploaded_by=uploaded_by, org_id=award.org_id,
            pipeline_stage="identified", source="manual", originating_award_id=award.id,
            # Version 3.0 architecture upgrade, Phase 14 — previously accepted
            # by RenewalCreate but never written anywhere; see FOARecord.
            renewal_notes=data.get("notes"),
        )
        db.add(record)
        await db.flush()
        await db.refresh(record)
        return record
