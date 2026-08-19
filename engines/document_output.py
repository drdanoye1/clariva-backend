"""
Engine 9 — Document Output Engine
Produces TXT, DOCX, and PDF exports of completed proposals.

Version 3.0 upgrade, "Real File Storage" scope (Phase B; see docs/
Clariva_File_Storage_Scoping_Document.docx) — export() used to write the
generated file to the API dyno's local disk (tempfile.gettempdir()/
sbir_exports/) and hand back a relative download_url read back from that
same path by routers/documents.py::download_file(). That local-disk path
never survived a Heroku dyno restart/deploy and would 404 the moment
traffic landed on a different dyno than the one that generated the file —
a live reliability bug, not a missing feature. Every exporter below now
writes into an in-memory io.BytesIO buffer instead (the exact pattern
already proven in routers/budget_export.py's six export endpoints), and
export() uploads that buffer to Cloudflare R2 via storage.py, logs a
StoredFile row, and returns a real presigned download URL. Nothing is ever
written to local disk anymore.
"""

from __future__ import annotations

import io
import logging
from datetime import datetime
from typing import Any, Dict, List, Optional

from sqlalchemy.ext.asyncio import AsyncSession

import storage
from models.db_models import StoredFile, new_uuid

_log = logging.getLogger(__name__)


class ExportBlockedError(Exception):
    """Raised by export() when require_clean_export=True and
    ComplianceEngine finds at least one error-severity violation. Carries the
    full ComplianceReport so the router can return it verbatim in the 422
    response — the caller needs to know exactly what to fix, not just that
    something failed. Raised before any file is generated or uploaded."""
    def __init__(self, report: Any):
        self.report = report
        super().__init__(f"Export blocked: {len(report.violations)} compliance issue(s), "
                          f"including at least one error-severity violation.")


AI_DISCLAIMER = (
    "AI-GENERATED CONTENT DISCLAIMER: This document was produced using artificial "
    "intelligence and is provided as a DRAFT ONLY. It may contain errors, omissions, "
    "or inaccuracies. This content MUST be thoroughly reviewed, verified, and approved "
    "by a qualified professional — including a subject-matter expert, grant writer, "
    "compliance officer, or legal advisor — before submission to any funding agency. "
    "Clariva is a productivity and drafting tool, not a substitute for professional "
    "judgment. The submitting organization bears sole responsibility for the accuracy, "
    "compliance, and appropriateness of any content submitted under its name."
)


class DocumentOutputEngine:
    """
    Generates submission-ready documents in TXT, DOCX, and PDF formats.
    """

    # Content-Type per format, for the R2 object metadata and the
    # presigned URL's Content-Disposition header.
    _CONTENT_TYPES = {
        "txt": "text/plain",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "pdf": "application/pdf",
    }

    async def export(
        self,
        proposal: Any,
        sections: List[Any],
        fmt: str,
        db: AsyncSession,
        created_by: str,
        org_id: Optional[str] = None,
        include_scoring: bool = True,
        include_reviewer: bool = False,
        include_compliance: bool = True,
        generate_figures: bool = False,
        format_options: Any = None,   # ExportFormatOptions schema object or None
        require_clean_export: bool = False,
        brand_template_key: Optional[str] = None,
    ) -> Dict[str, Any]:
        """Route to format-specific exporter, then persist the result to R2
        (storage.py) and log a StoredFile row. Flushes but does not commit —
        same "engines flush, routers commit" convention as every other
        engine in this codebase (e.g. award_engine.py).

        Phase 5 export QA gate — ComplianceEngine.validate() always runs
        first; its report rides along on the return dict's "compliance_report"
        key either way. When require_clean_export=True and the report has any
        error-severity violation, raises ExportBlockedError before generating
        or uploading anything — a 402/422-style "fix this first" gate, not a
        silent partial export.

        Phase 13 — brand_template_key resolves the document JSON schema's
        `brandingProfile` field (CLARIVA-DOCGEN-SPEC-001 §3) the same way
        agencyProfile is resolved just below: an explicit key overrides,
        None falls back to the org's own Organization.default_brand_
        template_key, which itself falls back to the hardcoded
        clariva_standard system template — see
        BrandTemplateEngine.resolve(). Applies to both DOCX and PDF (PDF is
        printed from the same DOCX buffer, see _export_pdf_via_docx()); TXT
        has no visual formatting concept to brand, so it's untouched.
        """
        from utils.doc_utils import FormatOptions
        from engines.compliance_engine import ComplianceEngine
        from engines.agency_profile_engine import AgencyProfileEngine
        from engines.brand_template_engine import BrandTemplateEngine

        # Phase 7 — resolve the admin-editable AgencyProfile (falls back to
        # ComplianceEngine's hardcoded limits when no DB override exists) and
        # feed it into validate() as an override rather than a DB lookup
        # inside validate() itself, keeping validate() synchronous.
        resolved_profile = await AgencyProfileEngine().resolve(db, getattr(proposal, "agency", None) or "OTHER")
        compliance_report = ComplianceEngine().validate(
            proposal, sections,
            section_limits_override=resolved_profile.section_limits,
            total_limit_override=resolved_profile.total_page_limit,
        )
        if require_clean_export and not compliance_report.passed:
            raise ExportBlockedError(compliance_report)

        # Phase 13 — resolve effective branding. If no explicit key was
        # passed, look up the org's own chosen default first (rather than
        # letting resolve()'s own None-means-"clariva_standard" default
        # apply prematurely) so an org's default template actually takes
        # effect on every export without every caller having to know it.
        effective_brand_key = brand_template_key
        if effective_brand_key is None and org_id:
            from sqlalchemy import select as _select
            from models.db_models import Organization as _Organization
            org_row = (await db.execute(_select(_Organization.default_brand_template_key).where(_Organization.id == org_id))).scalar_one_or_none()
            effective_brand_key = org_row
        resolved_brand = await BrandTemplateEngine().resolve(db, effective_brand_key, org_id)
        logo_bytes = await self._fetch_logo_bytes(resolved_brand.logo_url) if resolved_brand.logo_url else None

        # Convert Pydantic schema → dataclass (None preserves all defaults)
        opts: FormatOptions | None = None
        if format_options is not None:
            opts = FormatOptions(
                font=format_options.font,
                font_pt=format_options.font_pt,
                alignment=format_options.alignment,
                margins_in=format_options.margins_in,
                page_num_position=format_options.page_num_position,
                cover_page_number=format_options.cover_page_number,
                page_break_h1=format_options.page_break_h1,
                section_numbering=format_options.section_numbering,
                space_after_pt=format_options.space_after_pt,
                space_before_h1_pt=format_options.space_before_h1_pt,
            )

        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        safe_title = "".join(c if c.isalnum() or c in "-_" else "_" for c in proposal.title[:40])
        filename = f"{safe_title}_{timestamp}.{fmt}"

        # Pre-generate figures for DOCX/PDF when requested
        figures: Dict[str, bytes] = {}
        if generate_figures and fmt in ("docx", "pdf"):
            figures = await self._generate_figures(proposal, sections)

        buffer = io.BytesIO()
        if fmt == "txt":
            self._export_txt(buffer, proposal, sections)
        elif fmt == "docx":
            self._export_docx(buffer, proposal, sections, figures=figures, opts=opts,
                               brand=resolved_brand, logo_bytes=logo_bytes)
        elif fmt == "pdf":
            await self._export_pdf_via_docx(buffer, proposal, sections, figures=figures, opts=opts,
                                             brand=resolved_brand, logo_bytes=logo_bytes)
        else:
            raise ValueError(f"Unsupported format: {fmt}")

        content = buffer.getvalue()
        content_type = self._CONTENT_TYPES.get(fmt, "application/octet-stream")

        storage_key = await storage.upload_file(org_id, "proposal_export", content, filename, content_type)
        stored_file = StoredFile(
            id=new_uuid(), org_id=org_id, object_type="proposal_export", object_id=proposal.id,
            storage_key=storage_key, original_filename=filename, content_type=content_type,
            size_bytes=len(content), checksum=storage.sha256_hex(content), created_by=created_by,
        )
        db.add(stored_file)
        await db.flush()
        download_url = await storage.get_download_url(storage_key, filename=filename)

        return {
            "download_url": download_url,
            "file_size": len(content),
            "exported_at": datetime.utcnow(),
            "compliance_report": compliance_report,
        }

    # ── Brand template logo fetch (Phase 13) ─────────────────────────────────

    async def _fetch_logo_bytes(self, logo_url: Optional[str]) -> Optional[bytes]:
        """Best-effort fetch of a brand template's logo image, for embedding
        into the DOCX/PDF export. Same "degrade gracefully, never fail the
        export outright" convention as every other best-effort call in this
        codebase (e.g. engines/image_gen.py's DALL-E image fetch) — any
        failure (bad URL, network error, non-2xx, timeout) returns None and
        the export simply proceeds without a logo, exactly as if no
        logo_url had been configured at all."""
        if not logo_url:
            return None
        import httpx
        try:
            async with httpx.AsyncClient(timeout=15) as http:
                r = await http.get(logo_url)
                r.raise_for_status()
                return r.content
        except Exception as exc:
            _log.warning("Brand template logo fetch failed for %r: %s", logo_url, exc)
            return None

    # ── Figure pre-generation ─────────────────────────────────────────────────

    async def _generate_figures(self, proposal: Any, sections: List[Any]) -> Dict[str, bytes]:
        """
        Scan all section content for [FIGURE N: ...] and [IMAGE N: ...] markers,
        generate an image for each unique caption, and return a caption→bytes dict.
        """
        import re
        from engines import image_gen
        from config import settings

        MARKER_RE = re.compile(
            r'\[(?:FIGURE|IMAGE)\s*\d*\s*:\s*(.*?)\]',
            re.IGNORECASE,
        )
        captions: List[str] = []
        seen: set = set()
        for sec in sections:
            if not sec.content:
                continue
            for m in MARKER_RE.finditer(sec.content):
                cap = m.group(1).strip()
                if cap and cap not in seen:
                    seen.add(cap)
                    captions.append(cap)

        figures: Dict[str, bytes] = {}
        for cap in captions:
            img = await image_gen.generate_figure(cap, proposal.title, settings.OPENAI_API_KEY)
            if img:
                figures[cap] = img

        return figures

    # ── TXT ───────────────────────────────────────────────────────────────────

    def _export_txt(self, buffer: io.BytesIO, proposal: Any, sections: List[Any]) -> None:
        import textwrap
        disclaimer_wrapped = "\n".join(textwrap.wrap(AI_DISCLAIMER, width=68))
        lines = [
            "=" * 70,
            proposal.title.upper(),
            f"Agency: {proposal.agency}  |  Phase: {proposal.phase.replace('_', ' ').title()}",
            f"Status: {proposal.status}  |  Version: {proposal.version}",
            "=" * 70,
            "",
            "⚠  " + "─" * 65,
            disclaimer_wrapped,
            "─" * 68,
            "",
        ]
        for section in sections:
            if section.content:
                lines += [
                    f"\n{'─' * 70}",
                    f"  {section.title.upper()}",
                    f"{'─' * 70}",
                    "",
                    section.content,
                    "",
                ]
        lines += [
            "=" * 70,
            f"Generated by Clariva Intelligent Grant Writing Platform™",
            f"Exported: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}",
            "=" * 70,
        ]

        buffer.write("\n".join(lines).encode("utf-8"))

    # ── DOCX ──────────────────────────────────────────────────────────────────

    def _export_docx(self, buffer: io.BytesIO, proposal: Any, sections: List[Any],
                     figures: Dict[str, bytes] = {}, opts=None,
                     brand: Any = None, logo_bytes: Optional[bytes] = None) -> None:
        """`brand` is a DocumentBrandTemplateResolvedOut (or None — every
        field below degrades to today's unbranded rendering when it is).
        Phase 13 additions, in document order: an optional logo image
        before the title, the title run colored to brand.primary_color
        when set, an optional header_text line under the metadata line,
        and an optional footer_text line added after add_page_numbers()
        (see utils/doc_utils.py::add_brand_footer_text's own docstring for
        why call order matters there)."""
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor, Inches
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from utils.doc_utils import (
                apply_federal_margins, add_federal_heading, add_body_para,
                add_page_numbers, add_brand_footer_text, hex_to_rgbcolor,
                render_content, _make_run, _para_spacing,
                FEDERAL_FONT, H1_PT,
            )

            font = opts.font if opts else FEDERAL_FONT

            doc = Document()
            apply_federal_margins(doc, opts=opts)

            # ── Brand logo (Phase 13) — rendered first so it appears above
            # the title in document flow. Best-effort: a truncated/corrupt
            # image (logo_bytes fetched but not a real image format
            # python-docx/Pillow can decode) is caught and skipped rather
            # than failing the whole export.
            if logo_bytes:
                try:
                    logo_para = doc.add_paragraph()
                    logo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _para_spacing(logo_para, after=6, before=0)
                    logo_para.add_run().add_picture(io.BytesIO(logo_bytes), width=Inches(1.5))
                except Exception:
                    _log.warning("Brand logo image could not be embedded — skipping.", exc_info=True)

            # ── Title ─────────────────────────────────────────────────────────
            title_para = doc.add_paragraph()
            title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _para_spacing(title_para, after=4, before=0)
            title_run = _make_run(title_para, proposal.title, bold=True, pt=H1_PT, font=font)
            brand_color = hex_to_rgbcolor(getattr(brand, "primary_color", None)) if brand else None
            if brand_color is not None:
                title_run.font.color.rgb = brand_color

            # Metadata line
            meta = doc.add_paragraph()
            meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _para_spacing(meta, after=6)
            _make_run(meta,
                      f"Agency: {proposal.agency}  |  "
                      f"Phase: {proposal.phase.replace('_', ' ').title()}  |  "
                      f"Version: {proposal.version}",
                      pt=10, font=font)

            # ── Brand header text (Phase 13) — optional short line under the
            # metadata line, e.g. a partner co-brand name.
            if brand and getattr(brand, "header_text", None):
                header_para = doc.add_paragraph()
                header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _para_spacing(header_para, after=8)
                _make_run(header_para, brand.header_text, italic=True, pt=10, font=font)

            # ── AI Disclaimer box ─────────────────────────────────────────────
            disclaimer_para = doc.add_paragraph()
            _para_spacing(disclaimer_para, after=8)
            dr = disclaimer_para.add_run("AI-GENERATED DRAFT — REQUIRES PROFESSIONAL REVIEW\n")
            dr.bold = True; dr.font.size = Pt(9)
            dr.font.color.rgb = RGBColor(0x92, 0x40, 0x0E)
            br = disclaimer_para.add_run(AI_DISCLAIMER)
            br.font.size = Pt(9)
            br.font.color.rgb = RGBColor(0x44, 0x33, 0x00)
            pPr = disclaimer_para._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "FFF3CD")
            pPr.append(shd)
            pBdr = OxmlElement("w:pBdr")
            for side in ("top", "left", "bottom", "right"):
                bdr = OxmlElement(f"w:{side}")
                bdr.set(qn("w:val"), "single"); bdr.set(qn("w:sz"), "6")
                bdr.set(qn("w:space"), "4"); bdr.set(qn("w:color"), "D97706")
                pBdr.append(bdr)
            pPr.append(pBdr)
            # ─────────────────────────────────────────────────────────────────

            # ── Sections ──────────────────────────────────────────────────────
            # Word/PDF Report Generation Development Specification (CLARIVA-
            # DOCGEN-SPEC-001): sections with structured_content populated
            # render through the semantic block renderer (Phase 2). Since
            # Phase 4, proposal_generator.py::generate_section() populates
            # structured_content for every NEWLY generated section; sections
            # generated before Phase 4 shipped (or manually edited via the
            # raw-text PATCH endpoint, which clears structured_content) have
            # none and still go through the legacy render_content()
            # regex-parser — that fallback is permanent, not a stopgap, since
            # pre-existing/manually-edited content will never retroactively
            # gain block JSON. A structured render that raises for any reason
            # also falls back to render_content() rather than failing the
            # whole export — a bad/legacy blob must degrade gracefully, not
            # break someone's proposal export.
            from utils.block_renderer import render_structured_content

            for sec in sections:
                structured = getattr(sec, "structured_content", None)
                if not sec.content and not structured:
                    continue
                add_federal_heading(doc, sec.title, level=1, opts=opts)
                rendered = False
                if structured:
                    try:
                        render_structured_content(doc, structured, figures=figures, opts=opts)
                        rendered = True
                    except Exception:
                        _log.warning(
                            "Structured render failed for proposal section %r (id=%s) — "
                            "falling back to legacy render_content().",
                            getattr(sec, "section_id", "?"), getattr(sec, "id", "?"),
                            exc_info=True,
                        )
                if not rendered and sec.content:
                    render_content(doc, sec.content, figures=figures, opts=opts)

            # ── Page numbers ──────────────────────────────────────────────────
            add_page_numbers(doc, opts=opts)
            # ── Brand footer text (Phase 13) — MUST come after
            # add_page_numbers() so it renders below the page-number line.
            if brand and getattr(brand, "footer_text", None):
                add_brand_footer_text(doc, brand.footer_text, opts=opts)

            doc.save(buffer)

        except ImportError:
            self._export_txt(buffer, proposal, sections)

    # ── PDF ───────────────────────────────────────────────────────────────────

    async def _export_pdf_via_docx(
        self, buffer: io.BytesIO, proposal: Any, sections: List[Any],
        figures: Dict[str, bytes] = {}, opts=None,
        brand: Any = None, logo_bytes: Optional[bytes] = None,
    ) -> None:
        """Word/PDF Report Generation Development Specification
        (CLARIVA-DOCGEN-SPEC-001), Phase 8 — PDF is now "print the DOCX
        that's already correct," not a second renderer. Builds the exact
        same DOCX _export_docx() produces (full block-JSON structured
        render — figures, tables, schedules, callouts, references,
        everything Phases 1-7 built, plus Phase 13's logo/title-color/
        header/footer branding) into an in-memory buffer, then converts
        those bytes to PDF via headless LibreOffice (utils/pdf_convert.py).
        Falls back to the legacy reportlab _export_pdf() below — unbranded,
        kept permanently as a last resort — if no LibreOffice binary is
        available in this environment or the conversion fails for any
        reason, same "degrade gracefully, never fail the export outright"
        convention as every other exporter here. This is a deliberate, small
        scope reduction: the legacy reportlab path predates Phase 13 and is
        already a degraded fallback in every other respect (no structured
        blocks, no figures), so it staying unbranded too is consistent
        rather than a regression.
        """
        from utils import pdf_convert

        docx_buffer = io.BytesIO()
        self._export_docx(docx_buffer, proposal, sections, figures=figures, opts=opts,
                           brand=brand, logo_bytes=logo_bytes)
        docx_bytes = docx_buffer.getvalue()

        # _export_docx() itself silently falls back to TXT when python-docx
        # isn't importable (see its own except ImportError clause) — in that
        # case docx_bytes is actually plain text, not a real DOCX, and
        # handing it to soffice would either fail outright or produce a
        # garbage PDF. Detect that case up front via the DOCX/ZIP magic
        # bytes and skip straight to the legacy PDF path instead.
        if not docx_bytes.startswith(b"PK"):
            self._export_pdf(buffer, proposal, sections)
            return

        try:
            pdf_bytes = await pdf_convert.docx_bytes_to_pdf_bytes(docx_bytes)
            buffer.write(pdf_bytes)
            return
        except pdf_convert.LibreOfficeUnavailableError:
            _log.warning(
                "LibreOffice not available in this environment — falling back "
                "to the legacy reportlab PDF renderer for proposal %s.",
                getattr(proposal, "id", "?"),
            )
        except Exception:
            _log.warning(
                "LibreOffice DOCX->PDF conversion failed for proposal %s — "
                "falling back to the legacy reportlab PDF renderer.",
                getattr(proposal, "id", "?"), exc_info=True,
            )

        self._export_pdf(buffer, proposal, sections)

    def _export_pdf(self, buffer: io.BytesIO, proposal: Any, sections: List[Any]) -> None:
        """Legacy reportlab PDF renderer — reads only each section's flat
        `content` string (no figures, tables, schedules, callouts, or
        references), pre-dating the block-JSON structured render Phases
        1-7 built for DOCX. Since Phase 8, this is no longer the primary
        PDF path; it survives only as _export_pdf_via_docx()'s fallback
        for environments with no LibreOffice binary available. Not worth
        upgrading further — its whole reason to exist is as a dependency-
        free last resort, and any real capability gap should be closed by
        keeping the LibreOffice path working, not by extending this one."""
        try:
            from reportlab.lib.pagesizes import LETTER
            from reportlab.lib.styles import ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
            from reportlab.lib import colors
            from utils.doc_utils import get_pdf_styles, pdf_page_number, strip_for_pdf

            S = get_pdf_styles()
            NAVY = colors.HexColor("#1F4E79")

            disclaimer_style = ParagraphStyle(
                "DisclaimerStyle", fontName="Times-Roman", fontSize=9,
                leading=13, spaceAfter=6,
                backColor=colors.HexColor("#FFF3CD"),
                borderColor=colors.HexColor("#D97706"),
                borderWidth=1, borderPadding=8,
                textColor=colors.HexColor("#443300"),
            )
            disclaimer_label = ParagraphStyle(
                "DisclaimerLabel", fontName="Times-Bold", fontSize=9,
                textColor=colors.HexColor("#92400E"),
                backColor=colors.HexColor("#FFF3CD"),
                borderPadding=8,
            )

            doc_obj = SimpleDocTemplate(
                buffer, pagesize=LETTER,
                leftMargin=inch, rightMargin=inch,
                topMargin=inch, bottomMargin=inch,
            )

            story = [
                Paragraph(strip_for_pdf(proposal.title), S["title"]),
                Paragraph(
                    f"Agency: {proposal.agency}  |  "
                    f"Phase: {proposal.phase.replace('_', ' ').title()}",
                    S["small"],
                ),
                Spacer(1, 0.15 * inch),
                HRFlowable(width="100%", thickness=1, color=NAVY),
                Spacer(1, 0.08 * inch),
                Paragraph("AI-GENERATED DRAFT — REQUIRES PROFESSIONAL REVIEW", disclaimer_label),
                Paragraph(
                    AI_DISCLAIMER.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"),
                    disclaimer_style,
                ),
                Spacer(1, 0.15 * inch),
            ]

            for sec in sections:
                if not sec.content:
                    continue
                story.append(Paragraph(strip_for_pdf(sec.title), S["h1"]))
                for para_text in sec.content.split("\n\n"):
                    para_text = para_text.strip()
                    if not para_text:
                        continue
                    import re
                    if re.match(r'^#{1,3}\s', para_text) or re.match(r'^[A-Z]\.\s', para_text):
                        story.append(Paragraph(strip_for_pdf(para_text), S["h2"]))
                    else:
                        story.append(Paragraph(strip_for_pdf(para_text), S["body"]))
                story.append(Spacer(1, 0.12 * inch))

            story += [
                HRFlowable(width="100%", thickness=0.5, color=colors.grey),
                Spacer(1, 0.05 * inch),
                Paragraph(
                    f"<i>Generated by Clariva SBIR Intelligence Platform  |  "
                    f"{datetime.utcnow().strftime('%Y-%m-%d')}</i>",
                    S["small"],
                ),
            ]

            doc_obj.build(story,
                          onFirstPage=pdf_page_number,
                          onLaterPages=pdf_page_number)

        except ImportError:
            self._export_txt(buffer, proposal, sections)
