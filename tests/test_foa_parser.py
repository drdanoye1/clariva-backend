"""
Engine 1 — FOA Parsing Engine.
`parse()` calls OpenAI; everything tested here is the deterministic text
extraction and response-normalization logic around that call.
"""
from __future__ import annotations

import io

import docx as docx_lib
import pytest

from engines.foa_parser import FOAParserEngine


@pytest.fixture()
def engine():
    return FOAParserEngine()


# ── extract_text / extract_text_from_string ────────────────────────────────

def test_extract_text_from_string_strips_whitespace(engine):
    assert engine.extract_text_from_string("  hello world  \n") == "hello world"


def test_extract_text_plain_bytes_decodes_utf8(engine):
    text = engine.extract_text(b"Plain solicitation text", "solicitation.txt")
    assert text == "Plain solicitation text"


def test_extract_text_routes_docx_by_extension(engine):
    buf = io.BytesIO()
    doc = docx_lib.Document()
    doc.add_paragraph("Funding Opportunity Announcement")
    doc.add_paragraph("Deadline: March 1, 2027")
    doc.save(buf)

    text = engine.extract_text(buf.getvalue(), "foa.docx")
    assert "Funding Opportunity Announcement" in text
    assert "Deadline: March 1, 2027" in text


def test_extract_text_docx_includes_table_cells(engine):
    buf = io.BytesIO()
    doc = docx_lib.Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Phase I"
    table.rows[0].cells[1].text = "$275,000"
    doc.save(buf)

    text = engine.extract_text(buf.getvalue(), "budget_table.docx")
    assert "Phase I" in text and "$275,000" in text


# ── _parse_json (multi-strategy LLM JSON extraction) ────────────────────────

def test_parse_json_direct(engine):
    assert engine._parse_json('{"agency": "NSF"}') == {"agency": "NSF"}


def test_parse_json_strips_markdown_fence(engine):
    raw = '```json\n{"agency": "NIH", "phase": "phase_i"}\n```'
    assert engine._parse_json(raw) == {"agency": "NIH", "phase": "phase_i"}


def test_parse_json_extracts_embedded_object(engine):
    raw = 'Here is the result:\n{"agency": "DOE"}\nLet me know if you need more.'
    assert engine._parse_json(raw) == {"agency": "DOE"}


def test_parse_json_returns_empty_dict_on_garbage(engine):
    assert engine._parse_json("not json at all") == {}


# ── _normalize ───────────────────────────────────────────────────────────────

def test_normalize_fills_in_missing_defaults(engine):
    result = engine._normalize({})
    assert result["agency"] == "OTHER"
    assert result["program_title"] == "Unknown Program"
    assert len(result["ordered_sections"]) > 0


def test_normalize_preserves_provided_values(engine):
    result = engine._normalize({"agency": "NSF", "program_title": "SBIR Phase I"})
    assert result["agency"] == "NSF"
    assert result["program_title"] == "SBIR Phase I"


def test_normalize_renormalizes_weights_to_sum_to_one(engine):
    result = engine._normalize({"weights": {"a": 2, "b": 2}})
    assert result["weights"]["a"] == pytest.approx(0.5)
    assert result["weights"]["b"] == pytest.approx(0.5)


def test_normalize_leaves_empty_weights_alone(engine):
    result = engine._normalize({"weights": {}})
    assert result["weights"] == {}
