"""
Word/PDF Report Generation Development Specification (CLARIVA-DOCGEN-
SPEC-001), Phase 14 — "Visual regression corpus and profile-specific
golden files" (§14 P2 priority row; §13 acceptance criteria: "Automated
regression tests include representative NSF, NIH, DOE/DOD and generic/
private-foundation documents").

Pixel-level visual diffing (rendering a DOCX to page images and comparing
pixels byte-for-byte) is fragile across environments — font substitution,
hinting, and anti-aliasing differ between a developer's machine, this
sandbox, and the CI/Heroku build image, even when the exact same
LibreOffice binary produces the PDF. A pixel diff would fail on font-
rendering noise as often as it caught a real formatting regression,
defeating the point of a regression gate. Instead, extract_snapshot()
below pulls a deterministic STRUCTURAL snapshot out of a rendered .docx —
paragraph text/style/alignment/bold, table dimensions and cell text,
footer text and page-number field presence, embedded image count — and
diff_snapshots() compares that. This is exactly the information a human
reviewer would notice as "the formatting changed" (a heading demoted, a
table losing a column, page numbers disappearing), without being
sensitive to environment-specific pixel noise. Two renders of identical
logical content don't even produce byte-identical .docx files (Word XML
embeds revision ids/timestamps), so this format is deliberately engineered
around python-docx's parsed document model, not raw XML/bytes.

pdf_page_count_via_soffice() below is the module's nod to the spec's
"render-to-image checks" phrase — a cheap, environment-stable SECONDARY
signal (did the layout blow up to twice as many pages, or collapse to
none) alongside the structural snapshot, not a replacement for it. It
never raises: no LibreOffice binary, a conversion failure, or a timeout
all just return None so this optional signal degrades gracefully wherever
it can't run, exactly like every other best-effort external-tool call in
this codebase (e.g. engines/document_output.py::_fetch_logo_bytes()).

Deliberately does NOT import utils/pdf_convert.py's docx_bytes_to_pdf_bytes
— that module imports config.py, which requires pydantic-settings to be
installed just to import. This module has no pydantic/sqlalchemy
dependency at all (only python-docx, plus pypdf for the page-count path),
so the primary structural-snapshot regression gate keeps working even in
an environment where the full app dependency stack isn't installed —
including this project's own sandbox, where every other exporter/engine
in this codebase pulls in an uninstallable dependency somewhere in its
import chain.
"""
from __future__ import annotations

import io
from typing import Any, Dict, List, Optional


def extract_snapshot(docx_bytes: bytes) -> Dict[str, Any]:
    """Open *docx_bytes* (a real .docx file's raw bytes) and return a JSON-
    safe, deterministic structural snapshot. Never raises for a
    well-formed .docx; a corrupt/non-.docx input surfaces python-docx's
    own exception to the caller (a golden-file/CI failure SHOULD be loud
    here, not silently swallowed)."""
    from docx import Document

    doc = Document(io.BytesIO(docx_bytes))

    paragraphs: List[Dict[str, Any]] = []
    for p in doc.paragraphs:
        text = p.text
        if not text.strip():
            continue  # blank spacer paragraphs are layout noise, not content
        first_run = p.runs[0] if p.runs else None
        paragraphs.append({
            "text": text,
            "style": p.style.name if p.style else None,
            "alignment": str(p.alignment) if p.alignment is not None else None,
            "bold": bool(first_run.bold) if first_run is not None else None,
        })

    tables: List[Dict[str, Any]] = []
    for t in doc.tables:
        tables.append({
            "rows": len(t.rows),
            "cols": len(t.columns),
            "cells": [[cell.text for cell in row.cells] for row in t.rows],
        })

    image_count = sum(1 for rel in doc.part.rels.values() if "image" in rel.reltype)

    footers: List[Dict[str, Any]] = []
    for section in doc.sections:
        footer_texts = [p.text for p in section.footer.paragraphs if p.text.strip()]
        footers.append({
            "texts": footer_texts,
            "has_page_number_field": _footer_has_page_field(section),
        })

    return {
        "paragraph_count": len(paragraphs),
        "paragraphs": paragraphs,
        "table_count": len(tables),
        "tables": tables,
        "image_count": image_count,
        "footers": footers,
    }


def _footer_has_page_field(section) -> bool:
    """"Page X of Y" (utils/doc_utils.py::add_page_numbers()) is inserted
    as a real Word FIELD (PAGE/NUMPAGES instrText inside a fldChar run),
    not literal text — python-docx's Paragraph.text on that run returns
    "", so footer_texts above would never show it. Detecting it directly
    from the footer's XML instead means a regression that silently drops
    the page-number field (a real formatting defect a reviewer would
    absolutely notice) is still caught by the snapshot."""
    xml = section.footer._element.xml
    return "PAGE" in xml and "NUMPAGES" in xml


def diff_snapshots(golden: Dict[str, Any], candidate: Dict[str, Any]) -> List[str]:
    """Human-readable diff messages for a test-failure/CI message. Returns
    [] (falsy) when the two snapshots are equivalent — callers should
    `assert not diff_snapshots(golden, candidate), "\\n".join(diffs)`."""
    diffs: List[str] = []

    if golden.get("paragraph_count") != candidate.get("paragraph_count"):
        diffs.append(
            f"paragraph count changed: golden={golden.get('paragraph_count')} "
            f"candidate={candidate.get('paragraph_count')}"
        )
    g_paras = golden.get("paragraphs", [])
    c_paras = candidate.get("paragraphs", [])
    for i in range(min(len(g_paras), len(c_paras))):
        if g_paras[i] != c_paras[i]:
            diffs.append(f"paragraph[{i}] differs: golden={g_paras[i]!r} candidate={c_paras[i]!r}")

    if golden.get("table_count") != candidate.get("table_count"):
        diffs.append(
            f"table count changed: golden={golden.get('table_count')} "
            f"candidate={candidate.get('table_count')}"
        )
    g_tables = golden.get("tables", [])
    c_tables = candidate.get("tables", [])
    for i in range(min(len(g_tables), len(c_tables))):
        if g_tables[i] != c_tables[i]:
            diffs.append(f"table[{i}] differs: golden={g_tables[i]!r} candidate={c_tables[i]!r}")

    if golden.get("image_count") != candidate.get("image_count"):
        diffs.append(
            f"image count changed: golden={golden.get('image_count')} "
            f"candidate={candidate.get('image_count')}"
        )

    if golden.get("footers") != candidate.get("footers"):
        diffs.append(f"footers differ: golden={golden.get('footers')!r} candidate={candidate.get('footers')!r}")

    return diffs


def pdf_page_count_via_soffice(docx_bytes: bytes, timeout_seconds: float = 60.0) -> Optional[int]:
    """Best-effort secondary regression signal — see module docstring.
    Deliberately duplicates (rather than imports) utils/pdf_convert.py's
    minimal soffice-invocation logic, synchronously, to avoid pulling in
    config.py's pydantic-settings dependency for what is otherwise a
    dependency-light test utility. Returns None on any failure; never
    raises."""
    import os
    import shutil
    import subprocess
    import tempfile
    import uuid

    binary = shutil.which("soffice") or shutil.which("libreoffice")
    if not binary:
        return None

    try:
        with tempfile.TemporaryDirectory(prefix="clariva_pdf_check_") as tmpdir:
            docx_path = os.path.join(tmpdir, "input.docx")
            with open(docx_path, "wb") as f:
                f.write(docx_bytes)

            profile_dir = os.path.join(tempfile.gettempdir(), f"clariva_lo_check_{uuid.uuid4().hex}")
            try:
                proc = subprocess.run(
                    [
                        binary, f"-env:UserInstallation=file://{profile_dir}",
                        "--headless", "--norestore", "--nolockcheck", "--nodefault", "--nofirststartwizard",
                        "--convert-to", "pdf:writer_pdf_Export",
                        "--outdir", tmpdir, docx_path,
                    ],
                    capture_output=True, timeout=timeout_seconds,
                )
            finally:
                shutil.rmtree(profile_dir, ignore_errors=True)

            if proc.returncode != 0:
                return None

            pdf_path = os.path.join(tmpdir, "input.pdf")
            if not os.path.exists(pdf_path):
                return None

            from pypdf import PdfReader
            return len(PdfReader(pdf_path).pages)
    except Exception:
        return None
