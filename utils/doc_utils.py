"""
Shared document formatting utilities — federal submission standard.

Federal spec: Times New Roman 12pt, 1" margins, single spacing,
6pt after paragraphs, bottom-centre page numbers (Page X of Y),
three-level heading hierarchy (16 / 14 / 12 pt bold).
"""

from __future__ import annotations

import re
from typing import List, Tuple


# ─── Markdown stripping ───────────────────────────────────────────────────────

def strip_markdown(text: str) -> str:
    """Remove all markdown markers; return plain text."""
    if not text:
        return text
    text = re.sub(r'\*\*\*(.+?)\*\*\*', r'\1', text, flags=re.DOTALL)
    text = re.sub(r'\*\*(.+?)\*\*',     r'\1', text, flags=re.DOTALL)
    text = re.sub(r'\*(.+?)\*',         r'\1', text, flags=re.DOTALL)
    text = re.sub(r'___(.+?)___',       r'\1', text, flags=re.DOTALL)
    text = re.sub(r'__(.+?)__',         r'\1', text, flags=re.DOTALL)
    text = re.sub(r'_(.+?)_',           r'\1', text, flags=re.DOTALL)
    text = re.sub(r'^#{1,6}\s+', '', text, flags=re.MULTILINE)
    text = re.sub(r'^[-*_]{3,}\s*$', '', text, flags=re.MULTILINE)
    text = re.sub(r'`(.+?)`', r'\1', text)
    return text


def split_bold(text: str) -> List[Tuple[str, bool]]:
    """
    Split text on **bold** markers.
    Returns [(segment, is_bold), ...].
    """
    parts = re.split(r'\*\*(.+?)\*\*', text, flags=re.DOTALL)
    return [(p, i % 2 == 1) for i, p in enumerate(parts) if p]


# ─── python-docx helpers ──────────────────────────────────────────────────────

FEDERAL_FONT    = "Times New Roman"
BODY_PT         = 12
H1_PT, H2_PT, H3_PT = 16, 14, 12
MARGIN_IN       = 1.0
SPACE_AFTER_PT  = 6
SPACE_BEFORE_PT = 0


def apply_federal_margins(doc) -> None:
    """Set all section margins to 1 inch."""
    from docx.shared import Inches
    for sec in doc.sections:
        sec.top_margin    = Inches(MARGIN_IN)
        sec.bottom_margin = Inches(MARGIN_IN)
        sec.left_margin   = Inches(MARGIN_IN)
        sec.right_margin  = Inches(MARGIN_IN)


def _para_spacing(para, after: int = SPACE_AFTER_PT, before: int = SPACE_BEFORE_PT) -> None:
    """Single line spacing, 6pt after, 0pt before, widow/orphan on."""
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pf = para.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after  = Pt(after)
    pPr = para._p.get_or_add_pPr()
    sp = pPr.find(qn('w:spacing'))
    if sp is None:
        sp = OxmlElement('w:spacing')
        pPr.append(sp)
    sp.set(qn('w:line'),     '240')
    sp.set(qn('w:lineRule'), 'auto')
    wc = pPr.find(qn('w:widowControl'))
    if wc is None:
        wc = OxmlElement('w:widowControl')
        pPr.append(wc)
    wc.set(qn('w:val'), '1')


def _make_run(para, text: str, bold: bool = False, italic: bool = False,
              pt: int = BODY_PT):
    from docx.shared import Pt
    from docx.oxml.ns import qn
    run = para.add_run(text)
    run.bold   = bold
    run.italic = italic
    run.font.size = Pt(pt)
    # Force Times New Roman through rPr/rFonts so it overrides the style default
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        from docx.oxml import OxmlElement
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rFonts.set(qn(attr), FEDERAL_FONT)
    return run


def add_body_para(doc, text: str) -> None:
    """Add a body paragraph, converting **bold** to bold runs. Skips blank text."""
    text = text.strip()
    if not text:
        return
    text = re.sub(r'^#{1,6}\s+', '', text)   # strip stray heading markers
    para = doc.add_paragraph()
    _para_spacing(para)
    for segment, is_bold in split_bold(text):
        _make_run(para, segment, bold=is_bold)


def add_federal_heading(doc, text: str, level: int = 1) -> None:
    """Add a heading at the given level with Times New Roman and federal sizes."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    text = re.sub(r'^#{1,6}\s+', '', text.strip())
    text = strip_markdown(text)
    pt_map = {1: H1_PT, 2: H2_PT, 3: H3_PT}
    pt     = pt_map.get(level, H2_PT)
    try:
        para = doc.add_paragraph(style=f"Heading {min(level, 3)}")
    except Exception:
        para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.clear()
    _para_spacing(para, after=6, before=(12 if level == 1 else 6))
    _make_run(para, text, bold=True, pt=pt)


def add_page_numbers(doc) -> None:
    """Insert bottom-centre 'Page X of Y' into every section footer."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    def _field_run(para, instr: str):
        run = para.add_run()
        run.font.size = Pt(10)
        rPr = run._r.get_or_add_rPr()
        rFonts = OxmlElement('w:rFonts')
        for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
            rFonts.set(qn(attr), FEDERAL_FONT)
        rPr.insert(0, rFonts)
        begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
        itext = OxmlElement('w:instrText'); itext.set(qn('xml:space'), 'preserve')
        itext.text = f' {instr} '
        end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
        run._r.extend([begin, itext, end])

    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _make_run(p, "Page ", pt=10)
        _field_run(p, "PAGE")
        _make_run(p, " of ", pt=10)
        _field_run(p, "NUMPAGES")


def add_figure_placeholder(doc, caption: str, fig_num: int | str = 1) -> None:
    """
    Insert a styled figure placeholder box (dashed blue border, light blue fill)
    followed by an italic caption.  Reviewers replace the box with the actual figure.
    """
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    # ── Placeholder box ──────────────────────────────────────────────────
    box = doc.add_paragraph()
    box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _para_spacing(box, after=2, before=10)

    pPr = box._p.get_or_add_pPr()
    # Dashed blue border
    pBdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        bdr = OxmlElement(f"w:{side}")
        bdr.set(qn("w:val"), "dashSmallGap")
        bdr.set(qn("w:sz"), "12")
        bdr.set(qn("w:space"), "4")
        bdr.set(qn("w:color"), "2563EB")
        pBdr.append(bdr)
    pPr.append(pBdr)
    # Light blue fill
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "EFF6FF")
    pPr.append(shd)

    label = box.add_run(f"\n[ INSERT FIGURE {fig_num} HERE ]\n")
    label.bold = True
    label.font.size = Pt(10)
    label.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    # ── Caption ──────────────────────────────────────────────────────────
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _para_spacing(cap, after=10, before=2)
    cr = cap.add_run(f"Figure {fig_num}. {caption}")
    cr.italic = True
    cr.font.size = Pt(10)
    _set_run_font(cr)


def add_figure_image(doc, caption: str, fig_num, image_bytes: bytes) -> None:
    """
    Insert an actual figure image (PNG bytes) into the document,
    centred with an italic caption below.
    Falls back to add_figure_placeholder() if insertion fails.
    """
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import io as _io

    try:
        img_para = doc.add_paragraph()
        img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _para_spacing(img_para, after=2, before=10)
        img_para.add_run().add_picture(_io.BytesIO(image_bytes), width=Inches(5.5))
    except Exception:
        add_figure_placeholder(doc, caption, fig_num)
        return

    # Caption
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _para_spacing(cap, after=10, before=2)
    # Strip structured-marker prefix for display
    import re as _re
    display_caption = _re.sub(r'^(FLOW|IMAGE)\s*\|[^|]*\|\s*', '', caption, flags=_re.IGNORECASE).strip()
    if not display_caption:
        display_caption = caption
    cr = cap.add_run(f"Figure {fig_num}. {display_caption}")
    cr.italic = True
    cr.font.size = Pt(10)
    _set_run_font(cr)


def _set_run_font(run) -> None:
    """Helper: force Times New Roman on a run's rFonts element."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(attr), FEDERAL_FONT)


def add_gantt_table(doc, full_content: str, title: str = "Project Schedule") -> None:
    """
    Parse '(Months X-Y)' patterns from *full_content* and render a colour-coded
    Gantt table.  Projects ≤ 12 months use individual month columns; longer
    projects are grouped into quarters (Q1, Q2, …).

    Falls back to a figure placeholder if no phase data is found.
    """
    import math
    from docx.shared import Pt, Inches, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    # ── Parse phases ─────────────────────────────────────────────────────
    pattern = re.compile(
        r'(?P<name>[A-Za-z][^\n(]{4,70}?)'
        r'\s*\(\s*[Mm]onths?\s+(?P<s>\d+)\s*[-–—]\s*(?P<e>\d+)',
    )
    phases = []
    seen = set()
    for m in pattern.finditer(full_content):
        name  = re.sub(r'\s+', ' ', m.group("name")).strip().rstrip(", ")
        start = int(m.group("s"))
        end   = int(m.group("e"))
        key   = (name[:30], start, end)
        if key not in seen:
            seen.add(key)
            phases.append((name, start, end))

    if not phases:
        add_figure_placeholder(doc, f"{title} — Gantt Chart (auto-generated from milestone text)", "G")
        return

    total = max(e for _, _, e in phases)
    use_quarters = total > 12

    if use_quarters:
        n_cols = math.ceil(total / 3)
        col_labels = [f"Q{i+1}" for i in range(n_cols)]
        def is_active(ps, pe, ci):
            qs = ci * 3 + 1; qe = ci * 3 + 3
            return ps <= qe and pe >= qs
    else:
        n_cols = total
        col_labels = [str(m) for m in range(1, n_cols + 1)]
        def is_active(ps, pe, ci):
            return ps <= ci + 1 <= pe

    # ── Caption heading ───────────────────────────────────────────────────
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _para_spacing(cap, after=3, before=10)
    cr = cap.add_run(f"Figure G. {title}")
    cr.italic = True; cr.bold = True; cr.font.size = Pt(10)
    _set_run_font(cr)

    # ── Build table ───────────────────────────────────────────────────────
    t = doc.add_table(rows=len(phases) + 1, cols=n_cols + 1)
    t.style = "Table Grid"

    # Column widths
    usable_w = 6.5  # inches (1" margins each side)
    name_w   = Inches(2.2)
    col_w    = Inches((usable_w - 2.2) / n_cols)
    for i, cell in enumerate(t.columns[0].cells):
        cell.width = name_w
    for ci in range(1, n_cols + 1):
        for cell in t.columns[ci].cells:
            cell.width = col_w

    # Header row
    hcells = t.rows[0].cells
    hcells[0].text = "Phase / Task"
    for ci, lbl in enumerate(col_labels):
        hcells[ci + 1].text = lbl

    for cell in hcells:
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _para_spacing(cp, after=1, before=1)
        if cp.runs:
            r = cp.runs[0]; r.bold = True; r.font.size = Pt(8)
            _set_run_font(r)
        # Dark header fill
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "1F4E79")
        tcPr.append(shd)
        if cp.runs:
            cp.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Data rows
    ACTIVE_FILL = "DBEAFE"   # blue-100
    ALT_FILL    = "F0F9FF"   # lighter alt row

    for ri, (name, ps, pe) in enumerate(phases):
        row = t.rows[ri + 1]
        # Phase name cell
        nc = row.cells[0]
        nc.text = name[:45]
        np_ = nc.paragraphs[0]
        _para_spacing(np_, after=1, before=1)
        if np_.runs:
            np_.runs[0].font.size = Pt(8)
            _set_run_font(np_.runs[0])
        # Row fill
        row_fill = ALT_FILL if ri % 2 else "FFFFFF"

        for ci in range(n_cols):
            cell = row.cells[ci + 1]
            tc   = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd  = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
            active = is_active(ps, pe, ci)
            shd.set(qn("w:fill"), ACTIVE_FILL if active else row_fill)
            tcPr.append(shd)
            if active:
                cp = cell.paragraphs[0]
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _para_spacing(cp, after=0, before=0)
                r = cp.add_run("█")
                r.font.size = Pt(7)
                r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    doc.add_paragraph()  # spacer after Gantt


def render_content(doc, content: str, figures: dict = {}) -> None:
    """
    Parse AI-generated text (may contain markdown) into *doc* using
    federal formatting.  Handles ## headings, A. section labels,
    |table| rows, numbered items, **bold** inline,
    [FIGURE N: caption] / [IMAGE N: caption] markers (inserts real image if
    the caption key exists in *figures* dict, otherwise inserts a placeholder),
    and [GANTT: title] Gantt charts.

    figures: dict mapping caption string → PNG bytes (pre-generated by image_gen).
    """
    lines = content.split("\n")
    tbl_buf: list = []   # list of list[str]
    fig_counter = [0]    # mutable figure counter

    def flush_table():
        if not tbl_buf:
            return
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        t = doc.add_table(rows=len(tbl_buf), cols=len(tbl_buf[0]))
        t.style = "Table Grid"
        for ri, row_cells in enumerate(tbl_buf):
            for ci, cell_text in enumerate(row_cells):
                cell = t.rows[ri].cells[ci]
                cell.text = ""
                cp = cell.paragraphs[0]
                _make_run(cp, cell_text, bold=(ri == 0))
                if ci > 0:
                    cp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        doc.add_paragraph()
        tbl_buf.clear()

    for line in lines:
        s = line.strip()

        if s.startswith("|"):
            cells = [c.strip() for c in s.split("|") if c.strip()]
            if all(set(c) <= set("-: ") for c in cells):
                continue   # separator
            if not tbl_buf:
                tbl_buf.append(cells)
            else:
                padded = (cells + [""] * len(tbl_buf[0]))[:len(tbl_buf[0])]
                tbl_buf.append(padded)
            continue
        else:
            flush_table()

        if not s:
            continue

        # ── [FIGURE N: caption] ──────────────────────────────────────────
        fig_m = re.match(r'^\[FIGURE\s*(\d*)\s*:\s*(.*?)\]$', s, re.IGNORECASE)
        if fig_m:
            fig_counter[0] += 1
            cap = fig_m.group(2).strip()
            img_bytes = figures.get(cap) if figures else None
            if img_bytes:
                add_figure_image(doc, cap, fig_counter[0], img_bytes)
            else:
                add_figure_placeholder(doc, cap, fig_counter[0])
            continue

        # ── [GANTT: title] ───────────────────────────────────────────────
        gantt_m = re.match(r'^\[GANTT:\s*(.*?)\]$', s, re.IGNORECASE)
        if gantt_m:
            add_gantt_table(doc, content, gantt_m.group(1).strip())
            continue

        # ── [IMAGE: caption] (alias for FIGURE) ─────────────────────────
        img_m = re.match(r'^\[IMAGE\s*(\d*)\s*:\s*(.*?)\]$', s, re.IGNORECASE)
        if img_m:
            fig_counter[0] += 1
            cap = img_m.group(2).strip()
            img_bytes = figures.get(cap) if figures else None
            if img_bytes:
                add_figure_image(doc, cap, fig_counter[0], img_bytes)
            else:
                add_figure_placeholder(doc, cap, fig_counter[0])
            continue

        md_h = re.match(r'^(#{1,3})\s+(.*)', s)
        if md_h:
            add_federal_heading(doc, md_h.group(2), level=len(md_h.group(1)))
            continue

        if re.match(r'^[A-Z]\.\s', s):
            add_federal_heading(doc, s, level=2)
            continue

        num_m = re.match(r'^(\d+)\.\s+(.*)', s)
        if num_m:
            para = doc.add_paragraph()
            _para_spacing(para)
            _make_run(para, f"{num_m.group(1)}. ")
            for seg, bold in split_bold(num_m.group(2)):
                _make_run(para, seg, bold=bold)
            continue

        add_body_para(doc, s)

    flush_table()


# ─── ReportLab helpers ────────────────────────────────────────────────────────

def strip_for_pdf(text: str) -> str:
    """Strip markdown and escape ReportLab XML special characters."""
    text = strip_markdown(text)
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def get_pdf_styles():
    """
    ParagraphStyle dict using Times-Roman (federal standard).
    Keys: title, h1, h2, body, bold, small.
    """
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_JUSTIFY

    NAVY = colors.HexColor("#1F4E79")
    return {
        "title": ParagraphStyle("fed_title", fontName="Times-Bold", fontSize=16,
                                textColor=NAVY, alignment=TA_CENTER, spaceAfter=6),
        "h1":    ParagraphStyle("fed_h1",    fontName="Times-Bold", fontSize=14,
                                textColor=NAVY, spaceAfter=4, spaceBefore=12),
        "h2":    ParagraphStyle("fed_h2",    fontName="Times-Bold", fontSize=12,
                                textColor=NAVY, spaceAfter=3, spaceBefore=8),
        "body":  ParagraphStyle("fed_body",  fontName="Times-Roman", fontSize=12,
                                alignment=TA_JUSTIFY, spaceAfter=6, leading=16),
        "bold":  ParagraphStyle("fed_bold",  fontName="Times-Bold",  fontSize=12,
                                spaceAfter=4),
        "small": ParagraphStyle("fed_small", fontName="Times-Roman", fontSize=10,
                                alignment=TA_CENTER, spaceAfter=2),
    }


def pdf_page_number(canvas, doc):
    """ReportLab canvas callback — draws 'Page X' at bottom centre."""
    from reportlab.lib.units import inch
    canvas.saveState()
    canvas.setFont("Times-Roman", 10)
    canvas.drawCentredString(doc.pagesize[0] / 2.0, 0.5 * inch,
                             f"Page {canvas.getPageNumber()}")
    canvas.restoreState()
