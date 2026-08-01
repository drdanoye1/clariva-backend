"""
Budget Builder router — CRUD, totals, justification generation, file extraction.
"""

from __future__ import annotations

import io
import json
import re
import uuid
from typing import Any, Dict, Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
import openai

from config import settings
from database import get_db
from models.db_models import BudgetRecord, OrgContextDB, Proposal, User
from routers.auth import get_current_user

router = APIRouter()

AGENCY_CAPS: Dict[str, Dict[str, float]] = {
    "NSF":    {"Phase I": 275_000,  "Phase II": 1_000_000},
    "NIH":    {"Phase I": 314_681,  "Phase II": 2_000_000},
    "DOD":    {"Phase I": 250_000,  "Phase II": 1_750_000},
    "DARPA":  {"Phase I": 250_000,  "Phase II": 1_500_000},
    "DOE":    {"Phase I": 275_000,  "Phase II": 1_750_000},
    "NASA":   {"Phase I": 200_000,  "Phase II": 750_000},
    "ARPA-E": {"Phase I": 250_000,  "Phase II": 1_500_000},
}

_client: Optional[openai.AsyncOpenAI] = None

def _get_client() -> openai.AsyncOpenAI:
    global _client
    if _client is None:
        _client = openai.AsyncOpenAI(api_key=settings.OPENAI_API_KEY)
    return _client


def _ai_error(exc: Exception) -> HTTPException:
    """Convert any AI service exception to a safe, user-facing HTTPException."""
    import logging, traceback
    log = logging.getLogger(__name__)
    if isinstance(exc, openai.RateLimitError):
        log.error("AI service rate limit exceeded: %s", exc)
        return HTTPException(status_code=503, detail=(
            "The AI service is temporarily unavailable due to high demand. "
            "Please try again in a few minutes."
        ))
    if isinstance(exc, openai.AuthenticationError):
        log.critical("AI service authentication failure: %s", exc)
        return HTTPException(status_code=503, detail="The AI service is not configured correctly. Please contact support.")
    if isinstance(exc, openai.APIConnectionError):
        log.error("AI service connection error: %s", exc)
        return HTTPException(status_code=503, detail="The AI service is unreachable. Please try again in a moment.")
    if isinstance(exc, openai.APIStatusError):
        log.error("AI service API error %s: %s", exc.status_code, exc.message)
        return HTTPException(status_code=503, detail="The AI service returned an unexpected error. Please try again.")
    log.error("Unexpected AI error:\n%s", traceback.format_exc())
    return HTTPException(status_code=500, detail="An unexpected error occurred. Please try again or contact support.")


# ── Calculation ───────────────────────────────────────────────────────────────

def _calc_totals(body: Dict[str, Any]) -> Dict[str, float]:
    months = float(body.get("budget_months") or 12)
    personnel_salary = 0.0
    personnel_fringe = 0.0
    for p in body.get("personnel") or []:
        sal = float(p.get("annual_salary") or 0)
        fringe = float(p.get("fringe_rate") or 0) / 100.0
        effort = float(p.get("effort_pct") or 0) / 100.0
        cost = sal * effort * (months / 12.0)
        personnel_salary += cost
        personnel_fringe += cost * fringe

    consultant_total  = sum(float(c.get("rate_per_day") or 0) * float(c.get("days") or 0) for c in (body.get("consultants") or []))
    equipment_total   = sum(float(e.get("cost") or 0) for e in (body.get("equipment") or []))
    travel_total      = sum(float(t.get("trips") or 0) * float(t.get("people") or 1) * float(t.get("cost_per_trip") or 0) for t in (body.get("travel") or []))
    other_total       = sum(float(o.get("cost") or 0) for o in (body.get("other_direct") or []))
    subcontract_total = sum(float(s.get("total_cost") or 0) for s in (body.get("subcontracts") or []))

    total_direct = personnel_salary + personnel_fringe + consultant_total + equipment_total + travel_total + other_total + subcontract_total

    if (body.get("indirect_base") or "mtdc").lower() == "mtdc":
        sub_excess = sum(max(0.0, float(s.get("total_cost") or 0) - 25_000.0) for s in (body.get("subcontracts") or []))
        base_amount = max(0.0, total_direct - equipment_total - sub_excess)
    else:
        base_amount = total_direct

    indirect_rate  = float(body.get("indirect_rate") or 0) / 100.0
    total_indirect = base_amount * indirect_rate
    total_cost     = total_direct + total_indirect
    fee_rate       = float(body.get("fee_rate") or 0) / 100.0
    fee_amount     = total_cost * fee_rate
    total_requested = total_cost + fee_amount

    return {
        "personnel_salary":   round(personnel_salary,  2),
        "personnel_fringe":   round(personnel_fringe,  2),
        "consultant_total":   round(consultant_total,  2),
        "equipment_total":    round(equipment_total,   2),
        "travel_total":       round(travel_total,      2),
        "other_total":        round(other_total,       2),
        "subcontract_total":  round(subcontract_total, 2),
        "total_direct":       round(total_direct,      2),
        "indirect_base_amount": round(base_amount,     2),
        "total_indirect":     round(total_indirect,    2),
        "total_cost":         round(total_cost,        2),
        "fee_amount":         round(fee_amount,        2),
        "total_requested":    round(total_requested,   2),
    }


def _resolve_cap(agency: str, phase: str) -> float:
    caps = AGENCY_CAPS.get(agency, {})
    phase_key = phase.replace("_", " ").title()
    return caps.get(phase_key) or caps.get("Phase I") or 300_000.0


def _record_to_dict(rec: BudgetRecord, proposal: Proposal, totals: dict | None = None) -> dict:
    body = {
        "budget_months": rec.budget_months, "indirect_rate": rec.indirect_rate,
        "indirect_base": rec.indirect_base, "fee_rate": rec.fee_rate or 7.0,
        "personnel": rec.personnel or [], "consultants": rec.consultants or [],
        "equipment": rec.equipment or [], "travel": rec.travel or [],
        "other_direct": rec.other_direct or [], "subcontracts": rec.subcontracts or [],
    }
    t = totals or _calc_totals(body)
    cap = _resolve_cap(proposal.agency, proposal.phase)
    return {**body, **t, "agency_cap": cap, "over_budget": t["total_requested"] > cap,
            "updated_at": rec.updated_at.isoformat() if rec.updated_at else None}


async def _get_proposal(proposal_id: str, owner_id: str, db: AsyncSession) -> Proposal:
    r = await db.execute(select(Proposal).where(Proposal.id == proposal_id, Proposal.owner_id == owner_id))
    p = r.scalar_one_or_none()
    if not p:
        raise HTTPException(status_code=404, detail="Proposal not found")
    return p


# ── GET ───────────────────────────────────────────────────────────────────────

@router.get("/{proposal_id}")
async def get_budget(
    proposal_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    proposal = await _get_proposal(proposal_id, current_user.id, db)
    r = await db.execute(select(BudgetRecord).where(BudgetRecord.proposal_id == proposal_id))
    rec = r.scalar_one_or_none()
    if not rec:
        months = 24 if "ii" in proposal.phase.lower() else 12
        cap = _resolve_cap(proposal.agency, proposal.phase)
        return {"budget_months": months, "indirect_rate": 0.0, "indirect_base": "mtdc", "fee_rate": 7.0,
                "personnel": [], "consultants": [], "equipment": [], "travel": [], "other_direct": [], "subcontracts": [],
                "personnel_salary": 0.0, "personnel_fringe": 0.0, "consultant_total": 0.0, "equipment_total": 0.0,
                "travel_total": 0.0, "other_total": 0.0, "subcontract_total": 0.0, "total_direct": 0.0,
                "indirect_base_amount": 0.0, "total_indirect": 0.0, "total_cost": 0.0,
                "fee_amount": 0.0, "total_requested": 0.0,
                "agency_cap": cap, "over_budget": False, "updated_at": None}
    return _record_to_dict(rec, proposal)


# ── PUT ───────────────────────────────────────────────────────────────────────

@router.put("/{proposal_id}")
async def save_budget(
    proposal_id: str,
    body: Dict[str, Any],
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    import logging, traceback
    log = logging.getLogger(__name__)

    proposal = await _get_proposal(proposal_id, current_user.id, db)
    try:
        r = await db.execute(select(BudgetRecord).where(BudgetRecord.proposal_id == proposal_id))
        rec = r.scalar_one_or_none()
        if not rec:
            rec = BudgetRecord(id=str(uuid.uuid4()), proposal_id=proposal_id)
            db.add(rec)

        rec.budget_months = int(body.get("budget_months") or 12)
        rec.indirect_rate = float(body.get("indirect_rate") or 0)
        rec.indirect_base = body.get("indirect_base") or "mtdc"
        rec.fee_rate      = float(body.get("fee_rate") or 7.0)
        rec.personnel     = body.get("personnel") or []
        rec.consultants   = body.get("consultants") or []
        rec.equipment     = body.get("equipment") or []
        rec.travel        = body.get("travel") or []
        rec.other_direct  = body.get("other_direct") or []
        rec.subcontracts  = body.get("subcontracts") or []

        totals = _calc_totals({**body})
        rec.total_direct   = totals["total_direct"]
        rec.total_indirect = totals["total_indirect"]
        rec.total_cost     = totals["total_cost"]
        await db.flush()
        # Refresh so server-side columns (updated_at, etc.) are loaded
        # before _record_to_dict accesses them synchronously.
        await db.refresh(rec)
    except Exception as exc:
        log.error("Budget save error for proposal %s:\n%s", proposal_id, traceback.format_exc())
        raise HTTPException(
            status_code=500,
            detail=f"Budget save failed: {type(exc).__name__}: {exc}",
        )
    return _record_to_dict(rec, proposal, totals)


# ── Generate Justification ────────────────────────────────────────────────────

@router.post("/{proposal_id}/generate-justification")
async def generate_justification(
    proposal_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    proposal = await _get_proposal(proposal_id, current_user.id, db)
    r = await db.execute(select(BudgetRecord).where(BudgetRecord.proposal_id == proposal_id))
    rec = r.scalar_one_or_none()
    if not rec:
        raise HTTPException(status_code=404, detail="Save a budget first before generating justification.")

    # Load profile for org name and PI
    pr = await db.execute(select(OrgContextDB).where(OrgContextDB.user_id == current_user.id))
    ctx = pr.scalar_one_or_none()
    org_name = (ctx and ctx.organization_name) or "the applicant organization"

    body = {
        "budget_months": rec.budget_months, "indirect_rate": rec.indirect_rate,
        "indirect_base": rec.indirect_base, "fee_rate": rec.fee_rate or 7.0,
        "personnel": rec.personnel or [], "consultants": rec.consultants or [],
        "equipment": rec.equipment or [], "travel": rec.travel or [],
        "other_direct": rec.other_direct or [], "subcontracts": rec.subcontracts or [],
    }
    totals = _calc_totals(body)
    months = rec.budget_months or 12
    hours_per_year = 2080.0

    # Build personnel detail for prompt
    personnel_lines = []
    for p in (rec.personnel or []):
        sal = float(p.get("annual_salary") or 0)
        effort = float(p.get("effort_pct") or 0)
        fringe = float(p.get("fringe_rate") or 0)
        hourly = sal / hours_per_year if sal else 0
        hours = (effort / 100.0) * (months / 12.0) * hours_per_year
        sal_cost = sal * (effort / 100.0) * (months / 12.0)
        fringe_cost = sal_cost * (fringe / 100.0)
        personnel_lines.append(
            f"- {p.get('name','[Name]')} ({p.get('role','Personnel')}): "
            f"${sal:,.0f}/yr, ${hourly:.2f}/hr, {hours:.0f} hours, "
            f"salary cost ${sal_cost:,.0f}, fringe {fringe}% = ${fringe_cost:,.0f}, "
            f"total ${sal_cost+fringe_cost:,.0f}"
        )

    consultant_lines = [
        f"- {c.get('name','Consultant')} ({c.get('role','')}): ${float(c.get('rate_per_day',0)):.0f}/day × {c.get('days',0)} days = ${float(c.get('rate_per_day',0))*float(c.get('days',0)):,.0f}"
        for c in (rec.consultants or [])
    ]
    equipment_lines = [f"- {e.get('name','')}: ${float(e.get('cost',0)):,.0f} — {e.get('description','')}" for e in (rec.equipment or [])]
    travel_lines = [
        f"- {t.get('purpose','')}, {t.get('destination','')}: {t.get('trips',0)} trip(s) × {t.get('people',1)} person(s) × ${float(t.get('cost_per_trip',0)):,.0f} = ${float(t.get('trips',0))*float(t.get('people',1))*float(t.get('cost_per_trip',0)):,.0f}"
        for t in (rec.travel or [])
    ]
    other_lines = [f"- {o.get('category','')}: {o.get('description','')} — ${float(o.get('cost',0)):,.0f}" for o in (rec.other_direct or [])]
    sub_lines = [f"- {s.get('organization','')}, PI: {s.get('pi_name','')}: ${float(s.get('total_cost',0)):,.0f}" for s in (rec.subcontracts or [])]

    # Build grant-type-aware label for prompt
    _grant_label = getattr(proposal, "program_label", None) or proposal.grant_type or "grant"
    _grant_label = _grant_label.replace("_", " ").title() if not any(c in _grant_label for c in ["(", " "]) else _grant_label

    prompt = f"""Write a formal, professional Budget Justification narrative for the following grant proposal.

GRANT PROGRAM: {_grant_label}
PROPOSAL: {proposal.title or 'Grant Proposal'}
ORGANIZATION: {org_name}
AGENCY / FUNDER: {proposal.agency}
PHASE / PERIOD: {proposal.phase.replace('_',' ').title()}
RESEARCH FOCUS: {proposal.research_focus or 'Project activities as described'}
PERFORMANCE PERIOD: {months} months

=== BUDGET DATA ===

PERSONNEL ({len(rec.personnel or [])} people):
{chr(10).join(personnel_lines) if personnel_lines else 'None'}

CONSULTANTS:
{chr(10).join(consultant_lines) if consultant_lines else 'None requested.'}

EQUIPMENT:
{chr(10).join(equipment_lines) if equipment_lines else 'None requested.'}

TRAVEL:
{chr(10).join(travel_lines) if travel_lines else 'None requested.'}

OTHER DIRECT COSTS:
{chr(10).join(other_lines) if other_lines else 'None requested.'}

SUBCONTRACTS:
{chr(10).join(sub_lines) if sub_lines else 'None requested.'}

=== COMPUTED TOTALS ===
Personnel Salaries:        ${totals['personnel_salary']:>12,.2f}
Fringe Benefits:           ${totals['personnel_fringe']:>12,.2f}
Consultants:               ${totals['consultant_total']:>12,.2f}
Equipment:                 ${totals['equipment_total']:>12,.2f}
Travel:                    ${totals['travel_total']:>12,.2f}
Other Direct Costs:        ${totals['other_total']:>12,.2f}
Subcontracts:              ${totals['subcontract_total']:>12,.2f}
Total Direct Costs:        ${totals['total_direct']:>12,.2f}
Indirect ({rec.indirect_rate}% on {rec.indirect_base.upper()}): ${totals['total_indirect']:>12,.2f}
Total Direct + Indirect:   ${totals['total_cost']:>12,.2f}
Small Business Fee ({rec.fee_rate or 7}%):  ${totals['fee_amount']:>12,.2f}
TOTAL REQUESTED:           ${totals['total_requested']:>12,.2f}

=== INSTRUCTIONS ===
Write a complete, formal budget justification appropriate for the grant program listed above. Structure with these sections:
- Executive Summary (project description, scope, what IS and IS NOT included)
- A. Senior/Key Personnel (each person individually: role, rate, hours, cost, why necessary)
- B. Other Personnel (if none, state explicitly)
- C. Fringe Benefits (rate source, calculation, total)
- D. Permanent Equipment (itemized or state none with reason)
- E. Travel (itemized or state none with reason)
- F. Participant Support Costs (state if not applicable and why)
- G. Other Direct Costs (itemized: consultants, materials, subcontracts, or state none)
- H. Total Direct Costs
- I. Indirect Costs (rate basis, MTDC/TDC explanation, total)
- J. Total Direct and Indirect Costs
- K. Fee / Profit (if applicable to this grant type; omit if not allowed)
- L. Total Amount Requested
- Budget Summary Table (plain text table listing all categories and dollar amounts)

Write in formal grant language appropriate for {_grant_label}. Do NOT use SBIR-specific language unless this is actually an SBIR grant. State explicitly when categories are $0 and why. Be specific about hourly rates, hours, and cost calculations.

CRITICAL FORMATTING RULES — the output will be inserted directly into a federal submission document:
- Do NOT use any markdown formatting: no **, no *, no ##, no __, no backticks, no bullet hyphens
- Write section labels as plain text (e.g. "A. Senior/Key Personnel:" not "**A. Senior/Key Personnel:**")
- Use plain prose paragraphs. Use numbered sub-items (1. 2. 3.) only where a list genuinely aids clarity.
- Write the budget summary table using plain pipe-delimited rows: Category | Amount (the exporter will format it)"""

    try:
        response = await _get_client().chat.completions.create(
            model=settings.OPENAI_MODEL,
            messages=[
                {"role": "system", "content": f"You are an expert grant writer specializing in {_grant_label} budget justifications. Write formal, precise, compliant budget narratives appropriate to this specific grant program. Output plain prose only — absolutely no markdown formatting (no **, no *, no ##, no backticks). Section labels like 'A. Senior/Key Personnel:' should appear as plain text."},
                {"role": "user", "content": prompt},
            ],
            temperature=0.3,
            max_tokens=4000,
        )
    except Exception as exc:
        raise _ai_error(exc)
    justification = response.choices[0].message.content.strip()
    # Post-process: strip SBIR language for non-SBIR grants
    from engines.proposal_generator import _sanitize_sbir_content
    grant_type_val = getattr(proposal, "grant_type", None) or "federal_other"
    justification = _sanitize_sbir_content(justification, grant_type_val, _grant_label)
    return {"justification": justification, "totals": totals}


# ── Extract Budget from File ──────────────────────────────────────────────────

@router.post("/{proposal_id}/extract-budget")
async def extract_budget_from_file(
    proposal_id: str,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Extract budget line items from an uploaded Excel, Word, PDF, or text file."""
    await _get_proposal(proposal_id, current_user.id, db)

    content = await file.read()
    filename = file.filename or ""
    fname_lower = filename.lower()
    text = ""

    # ── Excel ────────────────────────────────────────────────────────────────
    if fname_lower.endswith((".xlsx", ".xls", ".xlsm")):
        try:
            import pandas as pd
            dfs = pd.read_excel(io.BytesIO(content), sheet_name=None, header=None)
            parts = []
            for sheet_name, df in dfs.items():
                parts.append(f"Sheet: {sheet_name}")
                parts.append(df.fillna("").to_string(index=False, header=False))
            text = "\n\n".join(parts)
        except Exception as e:
            raise HTTPException(status_code=422, detail=f"Could not read Excel file: {e}")

    # ── Word ─────────────────────────────────────────────────────────────────
    elif fname_lower.endswith((".docx", ".doc")):
        try:
            from docx import Document
            doc = Document(io.BytesIO(content))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    parts.append("\t".join(c.text.strip() for c in row.cells))
            text = "\n".join(parts)
        except Exception as e:
            raise HTTPException(status_code=422, detail=f"Could not read Word file: {e}")

    # ── PDF ──────────────────────────────────────────────────────────────────
    elif fname_lower.endswith(".pdf"):
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                text = "\n\n".join(page.extract_text() or "" for page in pdf.pages[:20])
        except Exception:
            try:
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(content))
                text = "\n".join(page.extract_text() or "" for page in reader.pages[:20])
            except Exception as e:
                raise HTTPException(status_code=422, detail=f"Could not read PDF: {e}")

    # ── Plain text ───────────────────────────────────────────────────────────
    else:
        try:
            text = content.decode("utf-8", errors="replace")
        except Exception as e:
            raise HTTPException(status_code=422, detail=f"Could not read file: {e}")

    if not text.strip():
        raise HTTPException(status_code=422, detail="File appears to be empty or unreadable.")

    extract_prompt = f"""Extract budget line items from this document and return ONLY valid JSON matching this schema exactly:

{{
  "budget_months": integer,
  "indirect_rate": number,
  "indirect_base": "mtdc" or "tdc",
  "fee_rate": number,
  "personnel": [
    {{"name": string, "role": string, "annual_salary": number, "fringe_rate": number, "effort_pct": number}}
  ],
  "consultants": [
    {{"name": string, "role": string, "rate_per_day": number, "days": number}}
  ],
  "equipment": [
    {{"name": string, "description": string, "cost": number}}
  ],
  "travel": [
    {{"purpose": string, "destination": string, "trips": number, "people": number, "cost_per_trip": number}}
  ],
  "other_direct": [
    {{"category": string, "description": string, "cost": number}}
  ],
  "subcontracts": [
    {{"organization": string, "pi_name": string, "total_cost": number}}
  ],
  "extraction_notes": string
}}

Rules:
- If you see hourly rates and hours, convert to annual_salary = hourly_rate * 2080, effort_pct from hours
- If you see monthly salary, multiply by 12 for annual
- For fringe_rate: extract the percentage (e.g. "28%" → 28.0)
- For indirect_rate: extract the F&A or overhead rate percentage
- For fee_rate: look for "profit", "fee", or "award fee" percentage
- Return empty arrays [] for categories with no data
- Set extraction_notes to describe what you found and any assumptions made
- Return ONLY the JSON object, no markdown fences

DOCUMENT:
{text[:12000]}"""

    try:
        response = await _get_client().chat.completions.create(
            model=settings.OPENAI_MODEL,
            messages=[
                {"role": "system", "content": "You are an expert at extracting federal budget data from documents. Extract precise numbers."},
                {"role": "user", "content": extract_prompt},
            ],
            temperature=0.1,
            max_tokens=3000,
        )
    except Exception as exc:
        raise _ai_error(exc)

    raw = response.choices[0].message.content.strip()
    raw = re.sub(r"^```(?:json)?\s*", "", raw)
    raw = re.sub(r"\s*```$", "", raw)
    try:
        extracted = json.loads(raw)
    except json.JSONDecodeError:
        start, end = raw.find("{"), raw.rfind("}") + 1
        try:
            extracted = json.loads(raw[start:end])
        except Exception:
            raise HTTPException(status_code=500, detail="Could not parse extracted budget data.")

    # Add IDs to all array items
    import random, string
    def add_ids(items):
        return [{"id": ''.join(random.choices(string.ascii_lowercase, k=7)), **{k:v for k,v in item.items() if k != "id"}} for item in items]

    for field in ["personnel", "consultants", "equipment", "travel", "other_direct", "subcontracts"]:
        if field in extracted:
            extracted[field] = add_ids(extracted[field])

    totals = _calc_totals(extracted)
    return {"extracted": extracted, "totals": totals}
